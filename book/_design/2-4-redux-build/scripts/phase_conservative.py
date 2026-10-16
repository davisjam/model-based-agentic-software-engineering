"""Phase CONSERVATIVE (261015): the author's final eight-item revision pass.

No new slides, no new frameworks, no restructuring. The items:
  1. S23 — heading 'WHAT WAS MISSING?' -> 'WHAT MIGHT WE CHANGE?'; the five categories
     read as diagnostic alternatives ON the slide ('Not every failure requires all
     five.'); footer carries the absent-vs-inadequate distinction; unit mapping kept.
  2. S26 — the four-step FMEA PROCEDURE on the slide, above the worked table; the
     reveal staged in the notes; simplified-FMEA caveat in notes; RPN stays notes-only.
  3. S29 — 'Choose ONE model' moves into the VISIBLE task; FMEA question stays last.
  4. Timing — 5: 2->1, 9: 3->2, 17: 3->2 fund 23: 1->2, 24: 1->2, 25: 1->2 (author's
     preference order; slide 6's 4 min and slide 18's 4 min untouched). Sum stays 75.
  5. S13 — staging verified already present in notes (STAGING block); no change.
  6. S22 — the duplicate opening sentence is deleted; the Schoen line under the
     diagram (with the attribution) is the one kept.
  7. S28 — TRIGGER / AMPLIFIER labelled on the slide.
  8. S28 notes — what the step counts MEASURE; 23-vs-555 is illustration, not proof.
Plus: S8 notes qualify the 'weaknesses ALIGN' simplification (notes only).

Facts: VERIFIED-FACTS-261015.md — no new facts introduced by this pass.
Discipline: locate-and-assert-all -> mutate -> save -> verify (buildlib).
"""
from __future__ import annotations

import re
import shutil
import sys

sys.path.insert(0, "book/_design/2-4-redux-build/scripts")
from pptx.enum.text import PP_ALIGN
from pptx.util import Emu, Inches, Pt

from buildlib import (
    WIP, clone_into, delete_shape, fix_ids, load, one, place, retext,
    save_and_verify, set_notes, shapes_named, slide_title,
)

COURSE = ("course/lectures/act-2-delegation-and-control/"
          "04-failure-aware-engineering/2-4-Failure-Aware-Engineering.pptx")
TOL = Emu(Inches(0.02))


def notes_text(slide) -> str:
    return slide.notes_slide.notes_text_frame.text


def replace_note_line(slide, prefix: str, new_line: str) -> None:
    lines = notes_text(slide).split("\n")
    hits = [i for i, ln in enumerate(lines) if ln.startswith(prefix)]
    assert len(hits) == 1, f"note prefix {prefix!r}: {len(hits)} hits on {slide_title(slide)!r}"
    lines[hits[0]] = new_line
    set_notes(slide, "\n".join(lines))


def insert_note_after(slide, prefix: str, new_line: str) -> None:
    lines = notes_text(slide).split("\n")
    hits = [i for i, ln in enumerate(lines) if ln.startswith(prefix)]
    assert len(hits) == 1, f"note prefix {prefix!r}: {len(hits)} hits on {slide_title(slide)!r}"
    lines.insert(hits[0] + 1, new_line)
    set_notes(slide, "\n".join(lines))


def at(v_emu: int, inches: float) -> bool:
    return abs(v_emu - Inches(inches)) <= TOL


def shape_at(slide, base: str, left: float, top: float):
    hits = [sh for sh in shapes_named(slide, base) if at(sh.left, left) and at(sh.top, top)]
    assert len(hits) == 1, f"{base} at ({left},{top}): {len(hits)} hits on {slide_title(slide)!r}"
    return hits[0]


def set_all_runs(shape, pt: float | None = None, wrap: bool = True) -> None:
    shape.text_frame.word_wrap = wrap
    if pt is not None:
        for para in shape.text_frame.paragraphs:
            for run in para.runs:
                run.font.size = Pt(pt)


prs = load()
slides = list(prs.slides)
assert len(slides) == 30, len(slides)

# ============ LOCATE AND ASSERT ALL (no mutation yet) ============
EXPECT_TITLES = {
    5: "Explanation 1: A component-level failure",
    8: "Explanation 2: The Swiss cheese model",
    9: "The investigation identifies weaknesses",
    13: "What system are we investigating?",
    17: "Many computers do not imply independent failure",
    22: "Reflection changes the engineer",
    23: "Where will the next engineer encounter the lesson?",
    24: "Different lessons require different authorities",
    25: "How would we know the system is safer?",
    26: "Could we have anticipated this failure?",
    28: "What actually happened? July 2, 2019",
    29: "Diagnose the Cloudflare outage",
}
for idx, title in EXPECT_TITLES.items():
    got = slide_title(slides[idx - 1])
    assert got == title, f"slide {idx}: {got!r} != {title!r}"
s5, s8, s9, s13, s17 = (slides[i - 1] for i in (5, 8, 9, 13, 17))
s22, s23, s24, s25, s26, s28, s29 = (slides[i - 1] for i in (22, 23, 24, 25, 26, 28, 29))

# Item 5 — S13 staging must already be present in the notes (verify, change nothing).
assert "STAGING (no animation is authored in this deck" in notes_text(s13)
assert "Step 4" in notes_text(s13)

# S22: the duplicate intro label and the kept Schoen label.
s22_intro = [sh for sh in shapes_named(s22, "Label")
             if sh.has_text_frame and sh.text_frame.text.startswith("Engineers develop expertise")]
assert len(s22_intro) == 1
s22_schoen = [sh for sh in shapes_named(s22, "Label")
              if sh.has_text_frame and "Experience alone does not produce expertise" in sh.text_frame.text]
assert len(s22_schoen) == 1 and "Schön" in s22_schoen[0].text_frame.text

# S23: heading, grid rows, connectors, footer, mapping label.
s23_head = shape_at(s23, "ColumnHeading", 0.80, 1.14)
assert s23_head.text_frame.text == "WHAT WAS MISSING?"
ROW_TOPS_OLD = [1.45, 2.37, 3.29, 4.21, 5.13]
ROW_TOPS_NEW = [1.45, 2.27, 3.09, 3.91, 4.73]
s23_blocks = [shape_at(s23, "Block", 0.80, t) for t in ROW_TOPS_OLD]
s23_cards = [shape_at(s23, "Card", 4.20, t) for t in ROW_TOPS_OLD]
s23_conns = sorted(
    (sh for sh in s23.shapes if sh.name.startswith("Connector") and at(sh.left, 2.40)),
    key=lambda sh: sh.top)
assert len(s23_conns) == 4, len(s23_conns)
for conn, t in zip(s23_conns, [2.07, 2.99, 3.91, 4.83]):
    assert at(conn.top, t), (conn.top, t)
s23_footer = one(s23, "FooterBand")
assert s23_footer.text_frame.text.startswith("A failure is not a request for another gate")
s23_map = shape_at(s23, "Label", 0.60, 5.93)
assert s23_map.text_frame.text.startswith("obligation → Requirements")

# S26: lead label, headings, table rows, occurrence label.
s26_lead = shape_at(s26, "Label", 0.60, 0.92)
assert s26_lead.text_frame.text.startswith("The three models explain a failure")
s26_heads = [shape_at(s26, "ColumnHeading", x, 1.72) for x in (0.40, 2.30, 4.80, 6.90)]
S26_ROW_TOPS = [2.10, 2.95, 3.80, 4.65]
S26_SHIFT = 0.34
s26_rows = []
for t in S26_ROW_TOPS:
    s26_rows.append(shape_at(s26, "Block", 0.40, t))
    for x in (2.30, 4.80, 6.90):
        s26_rows.append(shape_at(s26, "Card", x, t))
s26_occ = shape_at(s26, "Label", 0.60, 5.58)
assert s26_occ.text_frame.text.startswith("How often will each mode occur?")

# S28: row blocks/captions, the step-count label to clone, the empty gap bands.
s28_rows = [shape_at(s28, "Block", 0.60, t) for t in (1.25, 2.30, 3.35, 4.40)]
assert [b.text_frame.text.split("\n")[0] for b in s28_rows] == [
    "CLOUDFLARE", "WAF MANAGED RULES", "THE DEFECTIVE RULE", "27 MINUTES"]
s28_step_label = shape_at(s28, "Label", 0.60, 5.03)
assert "23 steps" in s28_step_label.text_frame.text
# the gap bands (x>=3.70, under rows 2 and 3) must be empty of shapes
for sh in s28.shapes:
    if sh.left >= Inches(3.5):
        in_gap_2 = Inches(2.93) < sh.top < Inches(3.34)
        in_gap_3 = Inches(3.98) < sh.top < Inches(4.39)
        assert not (in_gap_2 or in_gap_3), f"gap occupied by {sh.name!r} at top={sh.top}"

# S29 footer.
s29_footer = one(s29, "FooterBand")
assert s29_footer.text_frame.text.startswith("YOUR TASK — choose an intervention")

# Notes lines that will be edited — assert each prefix exists exactly once.
for slide, prefix in [
    (s5, "TIMING: 2 minutes. FIRST OF THREE CAUSAL BOUNDARIES."),
    (s8, "NAMING + THE DISTINCTIVE CLAIM"),
    (s9, "TIMING: 3 minutes. THE EVIDENCE REPLACES THE SCHEMATIC."),
    (s17, "TIMING: 3 minutes. THE CENTRAL CONCEPTUAL PAYOFF OF MOVEMENT 3."),
    (s22, "EXPOSITION ORDER (author): read the intro sentence first"),
    (s23, "TIMING: 1 minutes. GOVERNANCE CONVERSION"),
    (s23, "PACING (author): the cover-the-right-column reveal is CUT"),
    (s23, "FACTUAL CARE:"),
    (s23, "EXPECTED ANSWERS AND MISCONCEPTION: 'add more gates.'"),
    (s24, "TIMING: 1 minutes. THE LOCAL AND SYSTEMIC ANALYSES MEET."),
    (s25, "TIMING: 1 minutes. CLOSES THE CROWDSTRIKE INVESTIGATION."),
    (s26, "ENGINEERING POINT:"),
    (s26, "RATINGS (notes only):"),
    (s28, "EXPOSITION ORDER (author):"),
    (s29, "ENGINEERING POINT:"),
]:
    hits = [ln for ln in notes_text(slide).split("\n") if ln.startswith(prefix)]
    assert len(hits) == 1, f"{prefix!r}: {len(hits)} hits on {slide_title(slide)!r}"

print("LOCATE-AND-ASSERT: all anchors found")

# ============ MUTATE ============

# ---- Item 1: slide 23 ----
retext(s23_head, "WHAT MIGHT WE CHANGE?")
for blk, card, t in zip(s23_blocks, s23_cards, ROW_TOPS_NEW):
    place(blk, top=t)
    place(card, top=t)
for conn, t in zip(s23_conns, [2.07, 2.89, 3.71, 4.53]):
    place(conn, top=t, height=0.20)
s23_diag = clone_into(s23, s23_map)
retext(s23_diag,
       "Diagnose the missing obligation, representation, evidence, evaluation, or enforcement.",
       "Not every failure requires all five.")
place(s23_diag, left=0.60, top=5.44, width=8.80, height=0.42)
set_all_runs(s23_diag, pt=12, wrap=True)
retext(s23_footer,
       "A failure is not a request for another gate.",
       "Absent and inadequate controls call for different responses.")
replace_note_line(
    s23, "TIMING: 1 minutes. GOVERNANCE CONVERSION",
    "TIMING: 2 minutes. GOVERNANCE CONVERSION — the lesson changes the ENVIRONMENT. (Reallocated +1 minute, "
    "author revision 261015: this slide introduces five diagnostic categories and maps them to four earlier "
    "course units — do not rush the walk.)")
replace_note_line(
    s23, "PACING (author): the cover-the-right-column reveal is CUT",
    "PACING (author): the cover-the-right-column reveal stays CUT — the FMEA exercise two slides ahead carries "
    "this movement's active learning. Spend the restored second minute on the five-row walk and the unit "
    "mapping, and state the slide-9 contrast yourself.")
replace_note_line(
    s23, "FACTUAL CARE:",
    "FACTUAL CARE: do NOT claim CrowdStrike lacked all five. Some controls existed and proved inadequate — the "
    "Content Validator RAN (against the wrong contract, RCA finding 4). A control that was ABSENT and one that "
    "ran against the WRONG OBLIGATION call for different responses: build the missing control versus re-point "
    "the existing one. The right-hand column is a proposed engineering response, not a description of "
    "CrowdStrike's actual remediation.")
replace_note_line(
    s23, "EXPECTED ANSWERS AND MISCONCEPTION: 'add more gates.'",
    "EXPECTED ANSWERS AND MISCONCEPTION: 'add more gates.' The footer is the rule: a failure is not a request "
    "for another gate — absent and inadequate controls call for different responses. Diagnose which category "
    "the chain broke at, then encode it; the on-slide line says it: not every failure requires all five.")

# ---- Item 2: slide 26 ----
retext(s26_lead,
       "The three models explain a failure that has happened. FMEA — Failure Mode and Effects Analysis —",
       "searches for failures that have not happened yet. The procedure:")
# Four steps as an aligned 2x2 grid (a centered two-line paragraph reads as ragged
# prose — visual-QA finding); each cell is its own left-aligned label.
PROC_STEPS = [
    (0.60, 1.52, "1  Identify a plausible failure mode"),
    (4.95, 1.52, "2  Determine its consequences"),
    (0.60, 1.78, "3  Identify existing or proposed controls"),
    (4.95, 1.78, "4  Decide whether the remaining risk warrants action"),
]
for x, y, txt in PROC_STEPS:
    cell = clone_into(s26, s26_occ)
    retext(cell, txt)
    place(cell, left=x, top=y, width=4.30, height=0.24)
    set_all_runs(cell, pt=12, wrap=False)
    for para in cell.text_frame.paragraphs:
        para.alignment = PP_ALIGN.LEFT
for sh in s26_heads:
    place(sh, top=1.72 + S26_SHIFT)
for sh in s26_rows:
    place(sh, top=Emu(sh.top).inches + S26_SHIFT)
place(s26_occ, top=5.86)
insert_note_after(
    s26, "ENGINEERING POINT:",
    "BUILD (ANIMATION-IN-SPEECH — no authored animation; run the reveal as staged speech): Step 1 — read the "
    "lead and the four-step PROCEDURE at the top, before any row: identify a plausible failure mode; determine "
    "its consequences; identify existing or proposed controls; decide whether the remaining risk warrants "
    "action. This is the method slide 27 asks students to run unaided. Step 2 — only then walk the worked "
    "table one row at a time, naming which procedure step each column instantiates (failure mode -> step 1, "
    "effect -> step 2, candidate control -> step 3). Step 3 — close on the occurrence line at the bottom: "
    "deciding which modes warrant action under unknown probabilities IS step 4.")
insert_note_after(
    s26, "RATINGS (notes only):",
    "SIMPLIFICATION (say it to the class): a full FMEA also assesses severity, occurrence, and detectability "
    "to prioritize action. We are using a simplified version because we lack reliable estimates.")

# ---- Item 3: slide 29 ----
retext(s29_footer,
       "YOUR TASK — choose ONE model. Explain the failure, propose an intervention,",
       "identify who can implement it, and specify evidence that it works.",
       "How should this incident change the FMEA for future WAF releases?")
insert_note_after(
    s29, "ENGINEERING POINT:",
    "VISIBLE TASK (author revision 261015): 'Choose ONE model' is now in the on-slide instructions — groups "
    "attempt one full analysis in the five minutes, not three shallow ones.")

# ---- Item 4: timing reallocation (5, 9, 17 fund 23, 24, 25; sum stays 75) ----
replace_note_line(
    s5, "TIMING: 2 minutes. FIRST OF THREE CAUSAL BOUNDARIES.",
    "TIMING: 1 minutes. FIRST OF THREE CAUSAL BOUNDARIES. DELIVER EFFICIENTLY (author reallocation 261015): "
    "reduce exposition — state the boundary, the model, and the intervention, then move; slide 6's 4-minute "
    "commitment discussion is protected and untouched.")
replace_note_line(
    s9, "TIMING: 3 minutes. THE EVIDENCE REPLACES THE SCHEMATIC.",
    "TIMING: 2 minutes. THE EVIDENCE REPLACES THE SCHEMATIC. DELIVER EFFICIENTLY (author reallocation 261015): "
    "slide 7 already introduced the evidence — walk the four rows and the absent/failed distinction without "
    "re-deriving each finding.")
replace_note_line(
    s17, "TIMING: 3 minutes. THE CENTRAL CONCEPTUAL PAYOFF OF MOVEMENT 3.",
    "TIMING: 2 minutes. THE CENTRAL CONCEPTUAL PAYOFF OF MOVEMENT 3. PACING (author reallocation 261015): "
    "state the common-cause principle once and do not re-derive it — slide 18 carries the architectural "
    "discussion and keeps its protected 4 minutes.")
replace_note_line(
    s24, "TIMING: 1 minutes. THE LOCAL AND SYSTEMIC ANALYSES MEET.",
    "TIMING: 2 minutes. THE LOCAL AND SYSTEMIC ANALYSES MEET. Pose the footer question and take one or two "
    "answers before stating the recognition yourself — students must reason about organizational authority, "
    "which needs a beat (author reallocation 261015).")
replace_note_line(
    s25, "TIMING: 1 minutes. CLOSES THE CROWDSTRIKE INVESTIGATION.",
    "TIMING: 2 minutes. CLOSES THE CROWDSTRIKE INVESTIGATION. Walk the rows with enough time for students to "
    "weigh what each evidence row establishes (author reallocation 261015) — the next slide re-invokes this "
    "slide's lesson (a control's existence is not evidence of its effectiveness) inside the FMEA's "
    "staged-rollout row.")

# ---- Item 6: slide 22 — keep the Schoen line, delete the duplicate opener ----
delete_shape(s22_intro[0])
replace_note_line(
    s22, "EXPOSITION ORDER (author): read the intro sentence first",
    "EXPOSITION ORDER (author, revised 261015): walk the chain, then land the distinction with the Schön line "
    "beneath the diagram — experience accumulates by itself; expertise does not. Reflection is the deliberate "
    "comparison of expectation with outcome, and that comparison, not the hours, feeds the repertoire. (The "
    "duplicate opening sentence above the diagram was removed: one statement of the point, stated once.)")

# ---- Items 7+8: slide 28 — label the two causal roles; notes on the step counts ----
s28_amp = clone_into(s28, s28_step_label)
retext(s28_amp, "AMPLIFIER — deployed globally in seconds, with no staged rollout")
place(s28_amp, left=3.70, top=2.96, width=5.70, height=0.26)
set_all_runs(s28_amp, pt=12, wrap=True)
s28_trig = clone_into(s28, s28_step_label)
retext(s28_trig, "TRIGGER — a pathological regex exhausts CPU resources")
place(s28_trig, left=3.70, top=4.01, width=5.70, height=0.26)
set_all_runs(s28_trig, pt=12, wrap=True)
insert_note_after(
    s28, "EXPOSITION ORDER (author):",
    "TRIGGER AND AMPLIFIER (now labelled on the slide): the pathological regex is the TRIGGER — it exhausts "
    "CPU; global deployment without staged containment is the AMPLIFIER — it set the blast radius. Both are "
    "needed to explain the incident's scale, and the pairing is the CrowdStrike parallel: there the content "
    "defect was the trigger and unstaged Rapid Response Content distribution the amplifier.")
insert_note_after(
    s28, "TRIGGER AND AMPLIFIER",
    "WHAT THE STEP COUNTS MEASURE (author): the counts are the number of candidate matching paths the engine "
    "explores before it can conclude — the work of backtracking — not merely a reflection of a longer input. "
    "And present 23-vs-555 as an ILLUSTRATION of the phenomenon, not a proof of exponential complexity: two "
    "measurements cannot establish an asymptotic growth rate.")

# ---- Slide 8 notes: qualify the 'weaknesses ALIGN' simplification (notes only) ----
insert_note_after(
    s8, "NAMING + THE DISTINCTIVE CLAIM",
    "QUALIFICATION (author, notes only — the slide wording stays): 'a failure escapes only when the weaknesses "
    "ALIGN' is a simplification. The intended meaning is that weaknesses align across the defenses RELEVANT TO "
    "A PARTICULAR FAILURE TRAJECTORY — not that every conceivable defense in the organization must fail. Carry "
    "the qualification verbally if a student presses on it.")

for s in (s23, s26, s28):
    fix_ids(s)

# ============ SAVE AND VERIFY ============
save_and_verify(prs, 30)

p2 = load()
# TIMING: per-slide expectations and the 75 total.
EXPECT_TIMING = [1, 2, 3, 4, 1, 4, 3, 2, 2, 2, 2, 1, 3, 3, 2, 4, 2, 4, 2, 3,
                 1, 1, 2, 2, 2, 3, 3, 3, 6, 2]
got = []
for i, s in enumerate(p2.slides, 1):
    m = re.search(r"TIMING:\s*(\d+)", s.notes_slide.notes_text_frame.text)
    assert m, f"slide {i} missing TIMING"
    got.append(int(m.group(1)))
assert got == EXPECT_TIMING, f"timing drift: {got}"
assert sum(got) == 75, sum(got)

# No-leak invariant: no Cloudflare incident detail on any slide BEFORE 28.
FORBIDDEN = ["backtrack", "27 minute", "13:42", "555", "cpu exhaustion",
             ".*(?:", "x=x", "14:0", "july 2, 2019", "2 july 2019"]
for i, s in enumerate(list(p2.slides)[:27], 1):
    body = []
    for sh in s.shapes:
        if sh.has_text_frame:
            for p in sh.text_frame.paragraphs:
                body.append("".join(r.text for r in p.runs))
        if sh.has_table:
            body.extend(c.text for row in sh.table.rows for c in row.cells)
    joined = "\n".join(body).lower()
    for tok in FORBIDDEN:
        assert tok not in joined, f"LEAK: {tok!r} on slide {i}"

# New on-slide content present.
def slide_text(s):
    return "\n".join(sh.text_frame.text for sh in s.shapes if sh.has_text_frame)
t23 = slide_text(list(p2.slides)[22])
assert "WHAT MIGHT WE CHANGE?" in t23 and "WHAT WAS MISSING?" not in t23
assert "Not every failure requires all five." in t23
assert "obligation → Requirements" in t23
assert "Absent and inadequate controls call for different responses." in t23
t26 = slide_text(list(p2.slides)[25])
assert "1  Identify a plausible failure mode" in t26
assert "4  Decide whether the remaining risk warrants action" in t26
assert "RPN" not in t26 and "Risk Priority" not in t26
t28 = slide_text(list(p2.slides)[27])
assert "TRIGGER — a pathological regex exhausts CPU resources" in t28
assert "AMPLIFIER — deployed globally in seconds, with no staged rollout" in t28
t29 = slide_text(list(p2.slides)[28])
assert "choose ONE model" in t29
assert "How should this incident change the FMEA for future WAF releases?" in t29
t22 = slide_text(list(p2.slides)[21])
assert "Experience alone does not produce expertise" in t22
assert "Engineers develop expertise by examining" not in t22

shutil.copyfile(WIP, COURSE)
print("INSTALLED: WIP copied to course path")
print("ALL PHASE-CONSERVATIVE CHECKS PASSED; TIMING sum =", sum(got))
