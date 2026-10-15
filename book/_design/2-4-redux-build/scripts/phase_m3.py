"""Phase M3: build Movement 3 (The system) — eight slides appended (45 -> 53).

Facts from VERIFIED-FACTS-261015.md §A. Mandatory edit 2 (safety constraint on
'The vendor is not the whole system'), correction to slide 13 (expanding ANALYTICAL
boundaries over a dependency/control graph; CrowdStrike and enterprise IT are separate
actors), recommendation 3 (compare architectures, not vendors).
"""
from __future__ import annotations

import sys

sys.path.insert(0, "book/_design/2-4-redux-build/scripts")
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_LINE_DASH_STYLE as MSO_LINE
from pptx.util import Inches, Pt

from buildlib import (
    clone_into, delete_shape, dup_slide, fix_ids, load, one, place, retext,
    save_and_verify, set_notes, shapes_named, slide_by_title,
)

RED = RGBColor(0xB3, 0x26, 0x1E)
GRAY = RGBColor(0x8A, 0x8A, 0x8A)


def set_line_geom(sh, x1, y1, x2, y2):
    """Position a line shape to run (x1,y1)->(x2,y2) with its arrowhead at (x2,y2)."""
    place(sh, min(x1, x2), min(y1, y2), abs(x2 - x1), abs(y2 - y1))
    xfrm = sh._element.spPr.xfrm
    if x2 < x1:
        xfrm.set("flipH", "1")
    else:
        xfrm.attrib.pop("flipH", None)
    if y2 < y1:
        xfrm.set("flipV", "1")
    else:
        xfrm.attrib.pop("flipV", None)


def set_size(sh, pt, line=None):
    paras = sh.text_frame.paragraphs if line is None else [sh.text_frame.paragraphs[line]]
    for para in paras:
        for r in para.runs:
            r.font.size = Pt(pt)


def make_boundary(slide, proto_block, left, top, width, height):
    b = clone_into(slide, proto_block)
    retext(b, "")
    place(b, left, top, width, height)
    b.fill.background()
    b.line.color.rgb = GRAY
    b.line.width = Pt(1.25)
    b.line.dash_style = MSO_LINE.DASH
    b.shadow.inherit = False
    return b


prs = load()
assert len(prs.slides._sldIdLst) == 45, "expected 45 slides at phase start"

t_stamp = slide_by_title(prs, "STAMP: Model the control structure")   # old 16
t_proto = slide_by_title(prs, "Explanation 1: A component-level failure")  # old 38
t_cards = slide_by_title(prs, "Preserve the example or preserve the lesson?")  # old 28
t_axis = slide_by_title(prs, "A broader model is not automatically a better model")  # old 17
proto_block = one(t_proto, "Block")
proto_card = one(t_proto, "Card")
proto_conn = one(t_proto, "Connector")  # horizontal line w/ arrow
proto_label = one(t_proto, "Label")
proto_footer = one(t_proto, "FooterBand")
# vertical arrow prototypes from the STAMP stack
stamp_conns = shapes_named(t_stamp, "Connector")
proto_down = next(c for c in stamp_conns if "flipV" not in c._element.spPr.xfrm.attrib and c.width == 0)
proto_up = next(c for c in stamp_conns if c._element.spPr.xfrm.get("flipV") == "1")


def stamp_dup_parts(s):
    """Index the duplicated STAMP slide's parts by their template text/geometry."""
    blocks = {b.text_frame.text.strip(): b for b in shapes_named(s, "Block")}
    conns = sorted([c for c in shapes_named(s, "Connector")],
                   key=lambda c: (c.top, c.left))
    labels = shapes_named(s, "Label")
    return blocks, conns, labels


# =============== Slide 13 — What system are we investigating? ===============
s = dup_slide(prs, t_proto)
retext(one(s, "Title"), "What system are we investigating?")
for sh in list(s.shapes):
    if sh.name.split(" ")[0] in ("Block", "Card", "Connector", "Label"):
        delete_shape(sh)
fb = one(s, "FooterBand")
retext(fb, "Which boundary contains the intervention we need?")

def blk(text2, l, t, w, h, lines):
    b = clone_into(s, proto_block)
    retext(b, *lines)
    place(b, l, t, w, h)
    return b

def crd(l, t, w, h, lines):
    c = clone_into(s, proto_card)
    retext(c, *lines)
    place(c, l, t, w, h)
    return c

blk("", 1.00, 1.45, 2.60, 0.60, ["CROWDSTRIKE ENGINEERING"])
blk("", 6.60, 1.45, 2.40, 0.60, ["ENTERPRISE IT"])
crd(1.30, 2.75, 1.40, 0.55, ["content"])
crd(3.10, 2.75, 1.60, 0.55, ["interpreter"])
crd(5.30, 2.75, 2.00, 0.55, ["Windows endpoint"])
crd(5.90, 4.05, 2.20, 0.55, ["critical services"])
for pts in [(2.30, 2.05, 2.00, 2.75),   # CrowdStrike -> content
            (2.70, 3.02, 3.10, 3.02),   # content -> interpreter
            (4.70, 3.02, 5.30, 3.02),   # interpreter -> endpoint
            (7.80, 2.05, 6.80, 2.75),   # enterprise -> endpoint
            (6.30, 3.30, 7.00, 4.05)]:  # endpoint -> services
    c = clone_into(s, proto_conn)
    set_line_geom(c, *pts)
# expanding analytical boundaries (separate actors; boundaries are models, not org charts)
for l, t, w, h, num in [(1.15, 2.58, 3.70, 0.92, "①"),
                        (0.90, 1.30, 4.20, 2.40, "②"),
                        (0.70, 1.15, 8.65, 3.65, "③"),
                        (0.55, 1.02, 8.95, 4.12, "④")]:
    make_boundary(s, proto_block, l, t, w, h)
    tag = clone_into(s, proto_label)
    retext(tag, num)
    place(tag, l + 0.03, t - 0.02, 0.4, 0.35)
    set_size(tag, 14)
extra = clone_into(s, proto_label)
retext(extra, "× thousands of organizations with the same dependency")
place(extra, 5.55, 4.70, 4.00, 0.35)
set_size(extra, 11)
legend = clone_into(s, proto_label)
retext(legend,
       "①  What caused the crash?  —  content + interpreter",
       "②  Why did the update escape?  —  CrowdStrike’s engineering system",
       "③  Why did organizations become unavailable?  —  deployment, endpoints, services",
       "④  Why could one failure disrupt many sectors?  —  the shared-dependency ecosystem")
place(legend, 0.60, 5.22, 8.90, 1.10)
set_size(legend, 12)
fix_ids(s)
set_notes(s, """TIMING: 3 minutes. OPENS MOVEMENT 3 — why could one failure affect so many?
ENGINEERING POINT: the system under investigation changes with the question. Each boundary is an ANALYTICAL choice — a model that preserves the relationships needed to answer one engineering question. Connect explicitly to the Modeling unit: a model preserves the relationships needed to answer an engineering question, and nothing else.
CARE: the boundaries are purposeful models, not containment relationships. CrowdStrike's engineering organization is NOT organizationally inside each customer's IT department — that is why CrowdStrike and enterprise IT appear as separate actors on one dependency/control graph, with boundaries drawn AROUND subsets of it.
WALK THE BOUNDARIES: ① content + interpreter — what caused the crash? ② CrowdStrike's engineering system — why did the update escape? ③ add enterprise deployment and the services it runs — why did organizations become unavailable? ④ the ecosystem of thousands of organizations holding the same dependency — why could one vendor's failure disrupt multiple sectors?
ASK: which boundary contains the intervention we need?
EXPECTED ANSWERS: 'the biggest one.' Misconception to correct: the larger boundary also exposes decisions OUTSIDE CrowdStrike's authority — breadth is bought with evidence and authority requirements, so you choose the boundary the decision needs.
TRANSITION: to reason about the larger boundaries we need a model built for systems — one whose vocabulary is control, not components.""")

# =============== Slide 14 — Explanation 3: Safety depends on control ===============
s = dup_slide(prs, t_stamp)
retext(one(s, "Title"), "Explanation 3: Safety depends on control")
blocks, conns, labels = stamp_dup_parts(s)
delete_shape(blocks["IAV"]); delete_shape(blocks["BOSCH"])
delete_shape(blocks["ORGANIZATION"]); delete_shape(blocks["ENGINEERING"])
# connectors: keep one down + one up pair, delete the rest (incl. side connectors)
keep_down, keep_up = None, None
for c in shapes_named(s, "Connector"):
    xfrm = c._element.spPr.xfrm
    vertical = c.width == 0
    if vertical and xfrm.get("flipV") != "1" and keep_down is None:
        keep_down = c
    elif vertical and xfrm.get("flipV") == "1" and keep_up is None:
        keep_up = c
    else:
        delete_shape(c)
ctrl = blocks["REGULATORS"]
retext(ctrl, "CONTROLLER", "process model: which actions are safe?")
place(ctrl, 3.00, 1.70, 4.00, 1.00)
set_size(ctrl, 12, line=1)
proc = blocks["OPERATING PROCESS"]
retext(proc, "CONTROLLED PROCESS")
place(proc, 3.00, 3.90, 4.00, 0.80)
set_line_geom(keep_down, 4.30, 2.70, 4.30, 3.90)
set_line_geom(keep_up, 5.70, 3.90, 5.70, 2.70)
lab_ctrl = next(l for l in labels if "control" in l.text_frame.text)
retext(lab_ctrl, "control actions ↓")
place(lab_ctrl, 1.20, 3.05, 2.60, 0.45)
lab_fb = next(l for l in labels if "feedback" in l.text_frame.text)
retext(lab_fb, "↑ feedback")
place(lab_fb, 7.10, 3.05, 2.30, 0.45)
lab_bottom = next(l for l in labels if "In this view" in l.text_frame.text)
retext(lab_bottom,
       "STAMP (Leveson): accidents arise when the control structure fails to enforce safety constraints.",
       "A controller can be software, an operator, an organization, or a regulator.")
place(lab_bottom, 0.60, 5.35, 8.80, 0.90)
fb = clone_into(s, proto_footer)
retext(fb, "What must the controller know and control?")
fix_ids(s)
set_notes(s, """TIMING: 4 minutes. THIRD CAUSAL BOUNDARY — introduce the model accurately.
ENGINEERING POINT: STAMP models safety as a CONTROL problem, not a component problem. Four elements: a controller issues control actions; a controlled process responds; feedback informs the controller; and the controller decides using a process model of the thing it controls. The system is safe when the control structure enforces its safety constraints.
CARE — three accuracy obligations: (1) STAMP is not a broader Swiss cheese, and not merely a big dependency diagram — its vocabulary is constraints, control, feedback, process models. (2) An unsafe outcome does not require a defective component: individually functioning components can interact to violate a system-level constraint. (3) Distinguish STAMP (the accident model) from STPA (the hazard-analysis method built on it) — we are introducing the model, not teaching the procedure.
ATTRIBUTION: Nancy Leveson, Engineering a Safer World (2011).
ASK: what must the controller know and control? Why can a controller with correct intentions still act unsafely?
EXPECTED ANSWERS: steer toward the three failure seats — an inaccurate process model, inadequate feedback, or insufficient control actions. Each is visible in this diagram, and none is a 'defective component.'
TRANSITION: apply it. What did CrowdStrike's update control structure look like on July 19?""")

# =============== Slide 15 — Who controlled the update? ===============
s = dup_slide(prs, t_stamp)
retext(one(s, "Title"), "Who controlled the update?")
blocks, conns, labels = stamp_dup_parts(s)
delete_shape(blocks["IAV"]); delete_shape(blocks["BOSCH"])
delete_shape(blocks["ORGANIZATION"])
downs, ups = [], []
for c in shapes_named(s, "Connector"):
    xfrm = c._element.spPr.xfrm
    if c.width != 0:
        delete_shape(c)
    elif xfrm.get("flipV") == "1":
        ups.append(c)
    else:
        downs.append(c)
for c in downs[2:]:
    delete_shape(c)
for c in ups[2:]:
    delete_shape(c)
downs, ups = downs[:2], ups[:2]
rel = blocks["REGULATORS"]
retext(rel, "CROWDSTRIKE RELEASE ORGANIZATION", "process model: is this update safe to distribute?")
place(rel, 2.70, 1.35, 4.60, 1.00)
set_size(rel, 12, line=1)
dist = blocks["ENGINEERING"]
retext(dist, "UPDATE DISTRIBUTION")
place(dist, 3.00, 2.95, 4.00, 0.70)
sens = blocks["OPERATING PROCESS"]
retext(sens, "FALCON SENSORS ON WINDOWS ENDPOINTS")
place(sens, 2.70, 4.25, 4.60, 0.80)
set_line_geom(downs[0], 4.30, 2.35, 4.30, 2.95)
set_line_geom(ups[0], 5.70, 2.95, 5.70, 2.35)
set_line_geom(downs[1], 4.30, 3.65, 4.30, 4.25)
set_line_geom(ups[1], 5.70, 4.25, 5.70, 3.65)
lab_ctrl = next(l for l in labels if "control" in l.text_frame.text)
retext(lab_ctrl, "release authorization ↓", "content ↓")
place(lab_ctrl, 0.80, 3.05, 2.40, 0.80)
lab_fb = next(l for l in labels if "feedback" in l.text_frame.text)
retext(lab_fb, "↑ validation results", "↑ endpoint telemetry")
place(lab_fb, 7.15, 3.05, 2.40, 0.80)
lab_bottom = next(l for l in labels if "In this view" in l.text_frame.text)
retext(lab_bottom, "CONSTRAINT — a content update must not disable the endpoint it protects.")
place(lab_bottom, 0.60, 5.55, 8.80, 0.50)
fb = clone_into(s, proto_footer)
retext(fb, "What did the release organization need to know before distributing?")
fix_ids(s)
set_notes(s, """TIMING: 3 minutes.
ENGINEERING POINT: a controller can make an unsafe decision because its process model is inaccurate, its feedback is inadequate, or its available control actions are insufficient — all three without anyone being negligent. For this release, the process model said 'the Content Validator passed it, so it is safe'; that model was wrong about what the validator actually checked (21 fields against the 20 supplied — RCA finding 4).
CARE: this is a SIMPLIFIED INSTRUCTIONAL control structure. Do not claim every depicted feedback path existed or operated during the incident, and do not invent telemetry, approvals, or rollout stages the RCA does not document. What the RCA does document: no staged rollout for Template Instances, and a remediation commitment to canary deployment rings with rollback and customer control over Rapid Response Content deployment (CrowdStrike RCA, 6 Aug 2024).
ASK: what did the release organization need to know before distributing the update?
EXPECTED ANSWERS: whether content satisfied the interpreter's assumptions; whether testing exercised the relevant behavior; whether telemetry could reveal danger quickly; whether distribution could be halted before exposure became widespread. Distinguish pre-release validation from operational feedback — both inform control, at different times.
TRANSITION: this structure stops at CrowdStrike's boundary. The constraint students actually care about — the hospital stays open — lives outside it.""")

# =============== Slide 16 — The vendor is not the whole system ===============
s = dup_slide(prs, t_stamp)
retext(one(s, "Title"), "The vendor is not the whole system")
blocks, conns, labels = stamp_dup_parts(s)
delete_shape(blocks["IAV"]); delete_shape(blocks["BOSCH"])
downs, ups = [], []
for c in shapes_named(s, "Connector"):
    xfrm = c._element.spPr.xfrm
    if c.width != 0:
        delete_shape(c)
    elif xfrm.get("flipV") == "1":
        ups.append(c)
    else:
        downs.append(c)
for c in downs[3:]:
    delete_shape(c)
for c in ups[1:]:
    delete_shape(c)
cs = blocks["REGULATORS"]
retext(cs, "CROWDSTRIKE", "release controls")
place(cs, 1.20, 1.50, 2.80, 0.90)
set_size(cs, 12, line=1)
ent = blocks["ORGANIZATION"]
retext(ent, "ENTERPRISE OPERATOR", "deployment + continuity controls")
place(ent, 5.80, 1.50, 3.00, 0.90)
set_size(ent, 12, line=1)
ep = blocks["ENGINEERING"]
retext(ep, "WINDOWS ENDPOINTS", "Falcon sensor + operating system")
place(ep, 2.80, 3.30, 4.40, 0.85)
set_size(ep, 12, line=1)
svc = blocks["OPERATING PROCESS"]
retext(svc, "CRITICAL ORGANIZATIONAL SERVICES", "clinical systems · aviation · finance")
place(svc, 2.80, 4.85, 4.40, 0.85)
set_size(svc, 12, line=1)
set_line_geom(downs[0], 2.60, 2.40, 3.40, 3.30)   # CrowdStrike -> endpoints
set_line_geom(downs[1], 7.30, 2.40, 6.50, 3.30)   # enterprise -> endpoints
set_line_geom(downs[2], 5.00, 4.15, 5.00, 4.85)   # endpoints support services
set_line_geom(ups[0], 8.75, 5.25, 8.75, 2.40)     # services/endpoints feedback to enterprise
lab_ctrl = next(l for l in labels if "control" in l.text_frame.text)
retext(lab_ctrl, "updates ↓")
place(lab_ctrl, 1.45, 2.55, 1.60, 0.40)
lab_fb = next(l for l in labels if "feedback" in l.text_frame.text)
retext(lab_fb, "availability,", "telemetry ↑")
place(lab_fb, 7.45, 3.60, 1.60, 0.70)
lab_bottom = next(l for l in labels if "In this view" in l.text_frame.text)
retext(lab_bottom,
       "SAFETY CONSTRAINT — failure of an endpoint security product",
       "must not disable essential organizational services.")
place(lab_bottom, 0.60, 5.90, 8.80, 0.80)
for para in lab_bottom.text_frame.paragraphs:
    for r in para.runs:
        r.font.bold = True
pol = clone_into(s, lab_ctrl)
retext(pol, "policies ↓")
place(pol, 6.05, 2.55, 1.30, 0.40)
fb = clone_into(s, proto_footer)
retext(fb, "Which controller can enforce this constraint?")
place(fb, 0.30, 6.78, 9.40, 0.55)
fix_ids(s)
set_notes(s, """TIMING: 4 minutes. MANDATORY PIVOT OF THE MOVEMENT — the constraint appears.
ENGINEERING POINT: state the system-level safety constraint explicitly: FAILURE OF AN ENDPOINT SECURITY PRODUCT MUST NOT DISABLE ESSENTIAL ORGANIZATIONAL SERVICES. Then ask which controller can enforce it. The answer is not CrowdStrike alone — the constraint cannot be satisfied entirely within the vendor's engineering system, and THAT is where STAMP earns its keep.
QUALIFICATION: this constraint is a CANDIDATE requirement for critical-service operators, not a documented universal obligation that existed before the outage.
FACTUAL BACKGROUND: CrowdStrike controls its software and content distribution. Customer organizations control aspects of deployment, continuity arrangements, and recovery. OS vendors control platform architecture. The responsibilities overlap but are not interchangeable: an enterprise cannot repair CrowdStrike's content-processing implementation; CrowdStrike cannot determine a hospital's continuity architecture.
ASK (2 minutes, pairs): which interventions require action by CrowdStrike, and which require action by its customers?
EXPECTED ANSWERS: bounds check / validator fix / staged rollout — vendor; deployment rings they can subscribe to, fallback procedures, continuity architecture — customer. Misconception: 'CrowdStrike should just fix it' — the constraint names a service property no vendor can deliver alone.
TRANSITION: why was the exposure so correlated in the first place? Name the failure class.""")

# =============== Slide 17 — Many computers do not imply independent failure ===============
s = dup_slide(prs, t_proto)
retext(one(s, "Title"), "Many computers do not imply independent failure")
for sh in list(s.shapes):
    if sh.name.split(" ")[0] in ("Block", "Card", "Connector", "Label"):
        delete_shape(sh)
retext(one(s, "FooterBand"), "Which apparently independent systems share a failure mode?")
orgs = ["HOSPITAL A", "AIRLINE B", "BANK C", "HOSPITAL D"]
xs = [0.70, 3.00, 5.30, 7.60]
for name, x in zip(orgs, xs):
    b = clone_into(s, proto_block)
    retext(b, name, "Falcon sensor")
    place(b, x, 1.50, 1.85, 0.90)
    set_size(b, 11, line=1)
src = clone_into(s, proto_block)
retext(src, "ONE DEFECTIVE", "CONTENT UPDATE")
place(src, 3.40, 4.90, 3.20, 0.95)
src.line.color.rgb = RED
src.line.width = Pt(2.0)
for x in xs:
    c = clone_into(s, proto_conn)
    set_line_geom(c, 5.00, 4.90, x + 0.925, 2.40)
    c.line.color.rgb = RED
mid = clone_into(s, proto_label)
retext(mid, "Redundancy within each organization did not help — every replica shared the dependency.")
place(mid, 0.60, 6.00, 8.80, 0.40)
fix_ids(s)
set_notes(s, """TIMING: 3 minutes. THE CENTRAL CONCEPTUAL PAYOFF OF MOVEMENT 3.
ENGINEERING POINT: common-cause failure — multiple components or organizations fail together because they share a dependency, condition, or mechanism. Redundancy provides limited protection when the redundant elements share the same vulnerability: a thousand machines running the same privileged component are not a thousand independent defenses against that component's failure.
FACTUAL SAFEGUARDS: the outage establishes widespread correlated EXPOSURE to CrowdStrike's software. It does NOT establish that every affected organization relied exclusively on CrowdStrike for endpoint security. Keep vendor concentration, shared software dependencies, and absence of redundancy distinct — related but different properties. Scale stays qualified: ~8.5M devices, under 1% of Windows machines (Microsoft, 20 Jul 2024).
ASK: which apparently independent systems in YOUR life share a failure mode? (Campus auth, cloud regions, package registries, certificate authorities.)
EXPECTED ANSWERS: students name redundant SERVERS. Push to the sharper question: the issue is not how many machines an organization runs, but how independently its essential functions can continue when a shared dependency fails.
TRANSITION: if shared dependency is the vulnerability, is the remedy to require diversity? Careful — compare architectures before prescribing vendors.""")

# =============== Slide 18 — How can critical services survive a shared failure? ===============
s = dup_slide(prs, t_cards)
retext(one(s, "Title"), "How can critical services survive a shared failure?")
cards = sorted(shapes_named(s, "Card"), key=lambda c: c.left)
assert len(cards) == 2
retext(cards[0], "SHARED FAILURE PATH")
place(cards[0], 0.60, 1.30, 4.30, 3.10)
cards[0].text_frame.paragraphs[0].runs[0].font.size = Pt(16)
retext(cards[1], "INDEPENDENT FAILURE PATHS")
place(cards[1], 5.10, 1.30, 4.30, 3.10)
cards[1].text_frame.paragraphs[0].runs[0].font.size = Pt(16)
# left mini-diagram: one vendor fans out to both services
va = clone_into(s, proto_block); retext(va, "Vendor A"); place(va, 1.95, 2.15, 1.60, 0.50)
s1 = clone_into(s, proto_card); retext(s1, "service 1"); place(s1, 0.95, 3.35, 1.55, 0.50)
s2 = clone_into(s, proto_card); retext(s2, "service 2"); place(s2, 3.00, 3.35, 1.55, 0.50)
c = clone_into(s, proto_conn); set_line_geom(c, 2.55, 2.65, 1.72, 3.35); c.line.color.rgb = RED
c = clone_into(s, proto_conn); set_line_geom(c, 2.95, 2.65, 3.78, 3.35); c.line.color.rgb = RED
cap = clone_into(s, proto_label); retext(cap, "one defect — both fail")
place(cap, 0.95, 3.95, 3.60, 0.35); set_size(cap, 11)
# right mini-diagram: primary path + independently-dependent fallback
vb = clone_into(s, proto_block); retext(vb, "Vendor A"); place(vb, 5.45, 2.15, 1.60, 0.50)
ind = clone_into(s, proto_block); retext(ind, "independent dependency"); place(ind, 7.35, 2.15, 1.90, 0.50)
set_size(ind, 10)
pp = clone_into(s, proto_card); retext(pp, "primary path"); place(pp, 5.45, 3.35, 1.60, 0.50)
fbk = clone_into(s, proto_card); retext(fbk, "fallback path"); place(fbk, 7.50, 3.35, 1.60, 0.50)
c = clone_into(s, proto_conn); set_line_geom(c, 6.25, 2.65, 6.25, 3.35); c.line.color.rgb = RED
c = clone_into(s, proto_conn); set_line_geom(c, 8.30, 2.65, 8.30, 3.35)
cap = clone_into(s, proto_label); retext(cap, "primary fails — service continues")
place(cap, 5.45, 3.95, 3.80, 0.35); set_size(cap, 11)
# requirement vs mechanisms
req = clone_into(s, proto_label)
retext(req,
       "REQUIREMENT — essential services remain available under specified endpoint failures",
       "CANDIDATE ARCHITECTURE — independent fallback paths · deployment isolation · delayed update rings",
       "CANDIDATE PROCUREMENT — vendor diversity (two vendors can still share an OS, a cloud, an admin plane)")
place(req, 0.60, 4.75, 9.00, 1.25)
set_size(req, 13)
fb = clone_into(s, proto_footer)
retext(fb, "What independence does the architecture actually need?")
fix_ids(s)
set_notes(s, """TIMING: 4 minutes. THE MAIN RISK SLIDE — do not let it become an argument for mandatory vendor diversity.
ENGINEERING POINT: compare ARCHITECTURES, not vendors. The question is whether critical functions possess independent failure paths. Separate three things: the REQUIREMENT (essential services remain available under specified endpoint failures), a candidate ARCHITECTURAL mechanism (independent fallback infrastructure, deployment isolation, delayed update rings), and a candidate PROCUREMENT mechanism (vendor diversity). The latter two are possible ways to satisfy the first — not conclusions that follow from the incident.
WHY VENDOR DIVERSITY IS NOT THE ANSWER KEY: two products may share an operating system, a cloud service, a configuration mechanism, or an administrative control plane — diversity of vendor names does not establish independence of failure modes. Conversely, an organization can obtain real resilience without a second endpoint product.
EPISTEMIC LADDER (keep it visible): OBSERVED — one vendor update produced widespread correlated failures. INFERRED — shared privileged dependencies created a common-cause vulnerability. PROPOSED — critical infrastructure might require more independence. UNRESOLVED — whether mandatory diversity would have helped, and whether its benefits justify its costs.
ASK (class discussion): what property should we require — multiple vendors, or continuity under a specified common-cause failure? Prescribing vendors constrains the solution; specifying the property leaves engineers free to satisfy it. Requirements vs Specification, exactly.
TRANSITION: whatever the requirement is — who has the authority to impose it?""")

# =============== Slide 19 — Who can impose a system-wide constraint? ===============
s = dup_slide(prs, t_stamp)
retext(one(s, "Title"), "Who can impose a system-wide constraint?")
blocks, conns, labels = stamp_dup_parts(s)
delete_shape(blocks["IAV"]); delete_shape(blocks["BOSCH"])
for c in shapes_named(s, "Connector"):
    if c.width != 0:
        delete_shape(c)
stack = [("REGULATORS", "REGULATORS / SECTOR AUTHORITIES"),
         ("ORGANIZATION", "CRITICAL-SERVICE OPERATORS"),
         ("ENGINEERING", "ENTERPRISE IT AND SECURITY"),
         ("OPERATING PROCESS", "ENDPOINTS AND SERVICES")]
for old, new in stack:
    retext(blocks[old], new)
lab_ctrl = next(l for l in labels if "control" in l.text_frame.text)
retext(lab_ctrl, "requirements,", "policies ↓")
place(lab_ctrl, 0.70, 2.60, 2.30, 0.80)
lab_fb = next(l for l in labels if "feedback" in l.text_frame.text)
retext(lab_fb, "↑ reports, evidence,", "operational feedback")
place(lab_fb, 7.00, 2.60, 2.40, 0.80)
lab_bottom = next(l for l in labels if "In this view" in l.text_frame.text)
retext(lab_bottom, "Not every sector has a regulator empowered to mandate these controls.")
fb = clone_into(s, proto_footer)
retext(fb, "Which controller has the authority to enforce independence?")
fix_ids(s)
set_notes(s, """TIMING: 3 minutes.
ENGINEERING POINT: controllers exist above the vendor, and the causal account may extend beyond the authority of the engineer conducting the investigation. A vendor can repair its software and release controls. A critical-service operator can change procurement, deployment, and continuity requirements. A sector authority may be able to impose resilience obligations across organizations — where its legal mandate permits. A cause outside your authority still belongs in the account; the response is to identify the responsible controller and ESCALATE the proposed intervention, not to omit the cause.
CARE: the diagram is conceptual. Do not imply every sector has a regulator empowered to mandate the proposed controls, and do not slide from 'a failure analysis reveals a systemic risk' to 'regulation is necessary' — necessity, effectiveness, and proportionality are further engineering and policy judgments requiring evidence.
ASK: which controller has the authority to enforce independence — and what feedback would it need to exercise that authority well?
EXPECTED ANSWERS: 'the government should require it.' Push back with the epistemic ladder from the last slide: proposed is not validated. Also surface the reverse misconception — 'if it's outside our authority it's out of scope.' Escalation is an engineering act.
TRANSITION: three models, three boundaries, three intervention scopes. Put them side by side.""")

# =============== Slide 20 — What did each model let us change? ===============
s = dup_slide(prs, t_axis)
retext(one(s, "Title"), "What did each model let us change?")
cards = sorted(shapes_named(s, "Card"), key=lambda c: c.left)
assert len(cards) == 4
texts = [c.text_frame.text.strip() for c in cards]
assert texts == ["Human error", "Swiss cheese", "Organizational", "STAMP"], texts
delete_shape(cards[2])  # Organizational
labels = shapes_named(s, "Label")
by_text = {l.text_frame.text.strip(): l for l in labels}
delete_shape(by_text["the shaping conditions"])
# reposition remaining three cards and their question labels
remaining = [cards[0], cards[1], cards[3]]
card_x = [0.80, 4.00, 7.20]
for c, x, t in zip(remaining, card_x, ["COMPONENT-LEVEL", "SWISS CHEESE", "STAMP"]):
    retext(c, t)
    place(c, x, 1.25, 2.00, 0.90)
qlabels = [by_text["the individual act"], by_text["the state of the defenses"],
           by_text["control and feedback loops"]]
for l, x, t in zip(qlabels, card_x, ["What failed?", "Why did it escape?", "Why so broad an impact?"]):
    retext(l, t)
    place(l, x - 0.10, 2.25, 2.20, 0.50)
# intervention row (cloned from the question labels)
for x, t in zip(card_x, ["repair the defect", "improve the defenses",
                         "change controls, dependencies, responsibilities"]):
    l = clone_into(s, qlabels[0])
    retext(l, t)
    place(l, x - 0.10, 2.85, 2.20, 0.75)
# push the axis below the two label rows; drop the redundant second axis row
conn = one(s, "Connector")
place(conn, 0.49, top=4.15)
for t, pos in [("Narrower", (0.49, 3.70)), ("Broader", (7.49, 3.70)),
               ("causal boundary", (3.83, 3.70))]:
    l = by_text[t]
    place(l, pos[0], pos[1])
for t in ["Less", "More", "structure"]:
    delete_shape(by_text[t])
big = by_text["Broader models expose more relationships.\nThey also demand more evidence and more authority to act."]
retext(big,
       "Broader explanations require more evidence and more authority to act on.",
       "Choose the boundary sufficient for the decision — broader is not automatically better.")
place(big, 0.93, 4.75, 8.20, 0.90)
fix_ids(s)
set_notes(s, """TIMING: 3 minutes. THE SYNTHESIS SLIDE — closes Movement 3.
ENGINEERING POINT: three models, three causal boundaries, three intervention scopes, encountered in the order the INVESTIGATION needed them. Component-level analysis asked 'what failed?' and let us repair the defect. Swiss cheese asked 'why did it escape?' and let us improve the defenses. STAMP asked 'why was the impact so broad?' and let us reason about controls, dependencies, and responsibilities — including controllers outside the vendor.
CARE: this is a progression of investigative SCOPE, not a ranking of theoretical sophistication, and not a ladder every analysis must climb. A local defect may justify a local correction; a defective release process may require stronger defenses; a systemic common-cause vulnerability may require architectural or institutional change. Do not suggest STAMP is the correct model for every failure. Swiss cheese and STAMP can each be applied at several scales — the boundary is the analyst's choice, not the theory's.
ASK: broader models expose more — what do they COST? (More evidence to support the account; more authority to act on it.)
EXPECTED ANSWERS AND MISCONCEPTION: students reflexively crown the broadest model. The analyst's obligation is a boundary SUFFICIENT for the decision at hand.
TRANSITION: we have candidate interventions at every scope. Which lessons does the evidence actually justify — and how do we keep them?""")

save_and_verify(prs, 53)
print("Movement 3 built: slides 46-53 appended")
