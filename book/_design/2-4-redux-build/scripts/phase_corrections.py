"""Phase CORRECTIONS (261016): author-approved correction pass — five substantive
fixes (slides 4, 5, 7, 17, 19), copyedits (6, 9, 20, 22, 28), optional slide-26
backtracking contrast, and the DOCUMENTED/INTERPRETATION/PROPOSED notes pass on 15-19.

28 slides in, 28 slides out. Facts from VERIFIED-FACTS-261015.md; the only derived
numbers are the 3- and 22-character input lengths, arithmetic on the quoted Cloudflare
strings ('x=x'; 'x=' followed by twenty 'x's).
"""
from __future__ import annotations

import sys

sys.path.insert(0, "book/_design/2-4-redux-build/scripts")
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_LINE_DASH_STYLE as MSO_LINE
from pptx.util import Pt

from buildlib import (
    clone_into, delete_shape, dup_slide, fix_ids, load, one, place, retext,
    save_and_verify, set_notes, shapes_named, slide_by_title,
)

RED = RGBColor(0xB3, 0x26, 0x1E)


def set_size(sh, pt, line=None):
    paras = sh.text_frame.paragraphs if line is None else [sh.text_frame.paragraphs[line]]
    for para in paras:
        for r in para.runs:
            r.font.size = Pt(pt)


def set_line_geom(sh, x1, y1, x2, y2):
    place(sh, min(x1, x2), min(y1, y2), abs(x2 - x1), abs(y2 - y1))
    xfrm = sh._element.spPr.xfrm
    if x2 < x1:
        xfrm.set("flipH", "1")
    else:
        xfrm.attrib.pop("flipH", None)
    if y2 < y1:
        xfrm.set("flipV", "1")
    else:
        xfrm.attrib.pop("flipV", None)


prs = load()
assert len(prs.slides._sldIdLst) == 28

s5 = slide_by_title(prs, "Explanation 1: A component-level failure")
proto_block = one(s5, "Block")
proto_card = one(s5, "Card")
proto_conn = one(s5, "Connector")
proto_label = one(s5, "Label")

# =============== FIX 1 — Slide 4: array-access diagram replaces the table ===============
s = slide_by_title(prs, "The content and the interpreter disagreed")
for sh in list(s.shapes):
    if sh.name.split(" ")[0] in ("Block", "RowCaption", "Connector"):
        delete_shape(sh)

def blk(lines, l, t, w, h):
    b = clone_into(s, proto_block); retext(b, *lines); place(b, l, t, w, h); return b

def crd(lines, l, t, w, h):
    c = clone_into(s, proto_card); retext(c, *lines); place(c, l, t, w, h); return c

def ln(x1, y1, x2, y2):
    c = clone_into(s, proto_conn); set_line_geom(c, x1, y1, x2, y2); return c

blk(["TEMPLATE DEFINITION: 21 FIELDS"], 3.00, 1.10, 4.00, 0.55)
crd(["content expects 21 inputs"], 0.90, 2.15, 3.60, 0.55)
crd(["integration code supplies 20"], 5.50, 2.15, 3.60, 0.55)
ln(4.40, 1.65, 2.70, 2.15)
ln(5.60, 1.65, 7.30, 2.15)
blk(["INPUT ARRAY [0 … 19]"], 3.40, 3.25, 3.20, 0.55)
ln(2.70, 2.70, 4.20, 3.25)
ln(7.30, 2.70, 5.80, 3.25)
crd(["interpreter attempts access [20]"], 3.20, 4.30, 3.60, 0.55)
ln(5.00, 3.80, 5.00, 4.30)
oob = blk(["OUT-OF-BOUNDS READ"], 3.50, 5.30, 3.00, 0.55)
oob.line.color.rgb = RED
oob.line.width = Pt(2.0)
ln(5.00, 4.85, 5.00, 5.30)
fix_ids(s)
notes = s.notes_slide.notes_text_frame
notes.text = notes.text + """
PACING (author): do not over-dwell on the technical mechanism — students need the INTERFACE MISMATCH, not the full Falcon content-processing architecture. Keep the walk-through under two minutes of the five.
TERMINOLOGY (stays in notes, off the slide): the Template Type defines the 21 fields; Template Instances are the content built against it; the Content Validator checked instances against the Template Type's 21 fields; the Content Interpreter consumed the integration code's 20-entry input array and, directed by the first non-wildcard 21st matching criterion, attempted the [20] access."""

# =============== FIX 2 — Slide 5: the defect is the incompatibility ===============
s = s5
b = one(s, "Block")
retext(b, "INCOMPATIBLE", "CONTENT + INTERPRETER")
set_size(b, 12, line=0)
set_size(b, 10, line=1)
cards = sorted(shapes_named(s, "Card"), key=lambda c: c.left)
assert cards[0].text_frame.text.strip() == "unsafe processing"
retext(cards[0], "unsafe access")
lab = one(s, "Label")
retext(lab, "Model: a chain of events traced to a defect.   Intervention: repair the incompatibility.")
notes = s.notes_slide.notes_text_frame
notes.text = notes.text + """
CORRECTION FOLDED IN: the first node reads 'incompatible content + interpreter', not 'defective content'. This was not a malformed file breaking an otherwise-correct interpreter — it was incompatible assumptions across the template definition, the integration code, the validator, and the interpreter. 'Repair the defect' therefore means repairing the incompatibility, not one content artifact; a content-only fix would leave the same mismatch armed for the next instance."""

# =============== FIX 5 — Slide 7: pipeline with attached weaknesses ===============
s = slide_by_title(prs, "How did the update reach production?")
blocks = sorted(shapes_named(s, "Block"), key=lambda b: b.top)
lqs = sorted(shapes_named(s, "LevelQuestion"), key=lambda q: q.top)
assert len(blocks) == 5 and len(lqs) == 5
from pptx.util import Emu
for b, q in zip(blocks, lqs):
    place(q, 4.75, width=4.95)
    yc = Emu(b.top).inches + Emu(b.height).inches / 2
    ln_s = clone_into(s, proto_conn)
    set_line_geom(ln_s, 4.00, yc, 4.75, yc)

# =============== FIX 3 — Slide 17: principle, not undocumented fact ===============
s = slide_by_title(prs, "Many computers do not imply independent failure")
lab = next(l for l in shapes_named(s, "Label") if "Redundancy" in l.text_frame.text)
retext(lab, "Replication does not provide independence when replicas share the same failure mode.")
set_size(lab, 14)

# =============== FIX 4 — Slide 19: explicit control + feedback, documented vs hypothetical ===============
s = slide_by_title(prs, "Who can impose a system-wide constraint?")
downs, ups = [], []
for c in shapes_named(s, "Connector"):
    if c._element.spPr.xfrm.get("flipV") == "1":
        ups.append(c)
    else:
        downs.append(c)
downs.sort(key=lambda c: c.top)
ups.sort(key=lambda c: c.top)
assert len(downs) == 3 and len(ups) == 3
downs[0].line.dash_style = MSO_LINE.DASH  # regulator-imposed resilience mandate = hypothetical
labels = shapes_named(s, "Label")
lab_ctrl = next(l for l in labels if "requirements" in l.text_frame.text)
lab_fb = next(l for l in labels if "reports" in l.text_frame.text)
retext(lab_ctrl, "resilience requirements ↓", "(hypothetical)")
place(lab_ctrl, 0.55, 1.90, 2.55, 0.65)
set_size(lab_ctrl, 12)
retext(lab_fb, "↑ incident reports, evidence")
place(lab_fb, 6.90, 1.95, 2.90, 0.40)
set_size(lab_fb, 12)
pairs = [("continuity policies ↓", 0.55, 3.00),
         ("deployment + recovery controls ↓", 0.55, 4.02),
         ("↑ service availability", 6.90, 3.05),
         ("↑ operational evidence, telemetry", 6.90, 4.07)]
for txt, x, y in pairs:
    l = clone_into(s, lab_fb)
    retext(l, txt)
    place(l, x, y, 2.55 if x < 5 else 2.90, 0.60)
    set_size(l, 12)
bot = next(l for l in labels if "Not every sector" in l.text_frame.text)
retext(bot,
       "Not every sector has a regulator empowered to mandate these controls.",
       "Dashed control = hypothetical addition to the structure, not a documented arrangement.")
place(bot, 0.60, 5.70, 8.80, 0.65)
set_size(bot, 13)
fix_ids(s)

# =============== COPYEDITS ===============
# Slide 9 — summary must cover all four findings
s = slide_by_title(prs, "The investigation identifies weaknesses")
retext(one(s, "FooterBand"),
       "Two defenses missing · one coverage gap · one ran against the wrong contract",
       "Source: CrowdStrike, External Technical Root Cause Analysis, 6 Aug 2024")

# Slide 6 — drop the title-duplicating block; enlarge the alternatives
s = slide_by_title(prs, "What would you change?")
delete_shape(one(s, "Block"))
for c in shapes_named(s, "Connector"):
    delete_shape(c)
cards = sorted(shapes_named(s, "Card"), key=lambda c: c.top)
assert len(cards) == 3
for c, y in zip(cards, [1.30, 2.85, 4.40]):
    place(c, 1.50, y, 7.00, 1.15)
notes = s.notes_slide.notes_text_frame
notes.text = notes.text.replace(
    "TIMING: 4 minutes. MAKE THEM COMMIT BEFORE THEY KNOW MORE.",
    "TIMING: 4 minutes — PROTECT 3 of them for the discussion (author pacing note). MAKE THEM COMMIT BEFORE THEY KNOW MORE.")

# Slide 20 — authority phrasing
s = slide_by_title(prs, "What did each model let us change?")
big = next(l for l in shapes_named(s, "Label") if "Broader explanations" in l.text_frame.text)
retext(big,
       "Broader explanations require more evidence and may require greater authority to act.",
       "Choose the boundary sufficient for the decision — broader is not automatically better.")
set_size(big, 15)

# Slide 22 — normalize stage-label capitalization
s = slide_by_title(prs, "Reflection changes the engineer")
for c in shapes_named(s, "Card"):
    t = c.text_frame.text.strip()
    if t in ("consequence", "repertoire"):
        retext(c, t.upper())

# Slide 28 — balance the two branches
s = slide_by_title(prs, "Failure should change future engineering")
blocks = {b.text_frame.text.split("\n")[0]: b for b in shapes_named(s, "Block")}
cards = {c.text_frame.text.strip(): c for c in shapes_named(s, "Card")}
place(blocks["REFLECTION"], 1.10, 4.55, 3.20, 0.85)
place(blocks["ENGINEERING STRUCTURE"], 5.70, 4.55, 3.20, 0.85)
set_size(blocks["REFLECTION"], 13, line=0)
set_size(blocks["ENGINEERING STRUCTURE"], 13, line=0)
place(cards["future judgment"], 1.10, 5.80, 3.20, 0.55)
place(cards["future constraints"], 5.70, 5.80, 3.20, 0.55)

# =============== OPTIONAL — Slide 26: backtracking work contrast ===============
s = slide_by_title(prs, "July 2, 2019: Cloudflare’s WAF outage")
la = clone_into(s, proto_label)
retext(la, "‘x=x’  →  23 steps")
place(la, 0.60, 5.03, 2.80, 0.30)
set_size(la, 12)
bar_a = clone_into(s, proto_block)
retext(bar_a, "")
place(bar_a, 3.60, 5.06, 0.26, 0.20)
lb = clone_into(s, proto_label)
retext(lb, "‘x=’ + twenty ‘x’s  →  555 steps")
place(lb, 0.60, 5.38, 2.80, 0.30)
set_size(lb, 12)
bar_b = clone_into(s, proto_block)
retext(bar_b, "")
place(bar_b, 3.60, 5.41, 6.00, 0.20)
bar_b.fill.solid()
bar_b.fill.fore_color.rgb = RED
cap = clone_into(s, proto_label)
retext(cap, "the same rule — the input grew from 3 to 22 characters")
place(cap, 3.60, 5.68, 5.60, 0.28)
set_size(cap, 10)
fix_ids(s)

# =============== NOTES PASS — slides 15-19: documented / interpretation / proposed ===============
set_notes(slide_by_title(prs, "Who controlled the update?"), """TIMING: 3 minutes.
ENGINEERING POINT: a controller can make an unsafe decision because its process model is inaccurate, its feedback is inadequate, or its available control actions are insufficient — all three without anyone being negligent.
DOCUMENTED (CrowdStrike RCA, 6 Aug 2024): the Content Validator validated Template Instances against the Template Type's 21 fields while the integration code supplied 20 (finding 4); Template Instances deployed to all endpoints with no staged rollout (finding 6); committed mitigations include canary deployment rings with rollback and customer control over Rapid Response Content deployment.
INTERPRETATION (ours, not the RCA's): the release organization's effective process model amounted to 'the Content Validator passed it, so it is safe to distribute' — and that model was wrong about what the validator actually checked. The RCA documents the validator defect; reading it as a controller's inaccurate process model is our STAMP-framed interpretation.
PROPOSED (instructional): the control structure drawn here is simplified. Do not claim every depicted feedback path existed or operated during the incident, and do not invent telemetry, approvals, or rollout stages the RCA does not document.
ASK: what did the release organization need to know before distributing the update?
EXPECTED ANSWERS: whether content satisfied the interpreter's assumptions; whether testing exercised the relevant behavior; whether telemetry could reveal danger quickly; whether distribution could be halted before exposure became widespread. Distinguish pre-release validation from operational feedback — both inform control, at different times.
TRANSITION: this structure stops at CrowdStrike's boundary. The constraint students actually care about — the hospital stays open — lives outside it.""")

set_notes(slide_by_title(prs, "The vendor is not the whole system"), """TIMING: 4 minutes. MANDATORY PIVOT OF THE MOVEMENT — the constraint appears.
ENGINEERING POINT: state the system-level safety constraint explicitly: FAILURE OF AN ENDPOINT SECURITY PRODUCT MUST NOT DISABLE ESSENTIAL ORGANIZATIONAL SERVICES. Then ask which controller can enforce it. The answer is not CrowdStrike alone — the constraint cannot be satisfied entirely within the vendor's engineering system, and THAT is where STAMP earns its keep.
DOCUMENTED: CrowdStrike controls its software and content distribution, and committed to giving customers control over Rapid Response Content deployment (RCA, 6 Aug 2024); affected organizations performed recovery across sectors (Microsoft, 20 Jul 2024). What any specific enterprise could or could not control is NOT documented in these sources.
INTERPRETATION (ours): the division of control drawn here — vendor release controls versus enterprise deployment and continuity controls — is a responsibility-allocation argument, not a finding of the postmortem. An enterprise cannot repair CrowdStrike's content-processing implementation; CrowdStrike cannot determine a hospital's continuity architecture.
PROPOSED (candidate requirement): the safety constraint itself. It is a candidate obligation for critical-service operators, not a documented universal obligation that existed before the outage.
ASK (2 minutes, pairs): which interventions require action by CrowdStrike, and which require action by its customers?
EXPECTED ANSWERS: bounds check / validator fix / staged rollout — vendor; deployment rings they can subscribe to, fallback procedures, continuity architecture — customer. Misconception: 'CrowdStrike should just fix it' — the constraint names a service property no vendor can deliver alone.
TRANSITION: why was the exposure so correlated in the first place? Name the failure class.""")

set_notes(slide_by_title(prs, "Many computers do not imply independent failure"), """TIMING: 3 minutes. THE CENTRAL CONCEPTUAL PAYOFF OF MOVEMENT 3.
ENGINEERING POINT: common-cause failure — multiple components or organizations fail together because they share a dependency, condition, or mechanism. The caption states the principle exactly: replication does not provide independence when replicas share the same failure mode. A thousand machines running the same privileged component are not a thousand independent defenses against that component's failure. That is an engineering principle, NOT a claim about any particular organization's redundancy architecture.
DOCUMENTED: ~8.5 million Windows devices affected — under 1% of all Windows machines (Microsoft, 20 Jul 2024); one defective content update was the shared trigger (CrowdStrike RCA, 6 Aug 2024).
INTERPRETATION (ours): the simultaneous multi-sector disruption is explained by correlated exposure to a shared privileged dependency. The incident does NOT establish that affected organizations had redundancy, that redundancy failed, or that any organization relied exclusively on CrowdStrike for endpoint security. Keep vendor concentration, shared software dependencies, and absence of redundancy distinct — related but different properties.
ASK: which apparently independent systems in YOUR life share a failure mode? (Campus auth, cloud regions, package registries, certificate authorities.)
EXPECTED ANSWERS: students name redundant SERVERS. Push to the sharper question: not how many machines an organization runs, but how independently its essential functions can continue when a shared dependency fails.
TRANSITION: if shared dependency is the vulnerability, is the remedy to require diversity? Careful — compare architectures before prescribing vendors.""")

set_notes(slide_by_title(prs, "How can critical services survive a shared failure?"), """TIMING: 4 minutes — PROTECT the discussion: 'what independence should critical services actually require?' deserves the time (author pacing note). THE MAIN RISK SLIDE — do not let it become an argument for mandatory vendor diversity.
ENGINEERING POINT: compare ARCHITECTURES, not vendors. The question is whether critical functions possess independent failure paths. Separate three things: the REQUIREMENT (essential services remain available under specified endpoint failures), a candidate ARCHITECTURAL mechanism (independent fallback infrastructure, deployment isolation, delayed update rings), and a candidate PROCUREMENT mechanism (vendor diversity). The latter two are possible ways to satisfy the first — not conclusions that follow from the incident.
DOCUMENTED: one vendor update produced widespread correlated failures (CrowdStrike RCA, 6 Aug 2024; Microsoft, 20 Jul 2024).
INTERPRETATION (ours): shared privileged dependencies created a common-cause vulnerability across organizations.
PROPOSED — and explicitly UNRESOLVED: critical infrastructure might require more independence; whether mandatory vendor diversity would have prevented this incident, and whether its benefits justify its costs, is unresolved. Two products may share an operating system, a cloud service, a configuration mechanism, or an administrative control plane — diversity of vendor names does not establish independence of failure modes. Conversely, an organization can obtain real resilience without a second endpoint product.
ASK (class discussion): what property should we require — multiple vendors, or continuity under a specified common-cause failure? Prescribing vendors constrains the solution; specifying the property leaves engineers free to satisfy it. Requirements vs Specification, exactly.
TRANSITION: whatever the requirement is — who has the authority to impose it?""")

set_notes(slide_by_title(prs, "Who can impose a system-wide constraint?"), """TIMING: 3 minutes.
ENGINEERING POINT: controllers exist above the vendor, and each level pairs a control action with the feedback it needs — read the diagram as control AND feedback, not as a hierarchy of rank. STAMP without feedback is just an org chart. A cause outside your authority still belongs in the account; the response is to identify the responsible controller and ESCALATE the proposed intervention, not to omit the cause.
DOCUMENTED vs HYPOTHETICAL (the dashed edge): enterprise deployment and recovery controls, and operators' continuity policies, are ordinary existing arrangements of the kind the incident exercised. A sector authority imposing resilience requirements for endpoint-security failure is a HYPOTHETICAL addition to the control structure — drawn dashed. Not every sector has a regulator empowered to mandate such controls, and no source in this lecture documents one that did so for this incident.
INTERPRETATION (ours): reading the outage as evidence of a systemic risk that some controller should manage is our interpretation of the documented exposure, not a finding of the RCA.
PROPOSED: any specific obligation at any level of this diagram. A failure analysis may reveal a systemic risk without establishing that regulation is necessary, effective, or proportionate — those are further engineering and policy judgments requiring evidence.
ASK: which controller has the authority to enforce independence — and what feedback would it need to exercise that authority well?
EXPECTED ANSWERS: 'the government should require it.' Push back with the epistemic ladder: proposed is not validated. Also surface the reverse misconception — 'if it's outside our authority it's out of scope.' Escalation is an engineering act.
TRANSITION: three models, three boundaries, three intervention scopes. Put them side by side.""")

save_and_verify(prs, 28)
print("corrections applied")
