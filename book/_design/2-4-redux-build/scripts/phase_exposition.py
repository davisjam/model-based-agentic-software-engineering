"""Phase EXPOSITION (261016): sequential-comprehensibility pass.

The single defect: the deck analyzed systems before introducing them. Fixes:
P0 establish-before-analyze (s3, s4, s14, s26); P1 model naming + sequencing
(s8, s13-notes, s21, s23); P2 prerequisite repairs (s7, s10, s19, s22);
s20 two-dimension correction (mechanism vs boundary); s1 notes; pacing rebalance
to a 75-minute sum. 28 slides in, 28 out; no new facts.
"""
from __future__ import annotations

import re
import sys

sys.path.insert(0, "book/_design/2-4-redux-build/scripts")
from pptx.dml.color import RGBColor
from pptx.util import Emu, Pt

from buildlib import (
    clone_into, delete_shape, fix_ids, load, one, place, retext,
    save_and_verify, shapes_named, slide_by_title,
)

RED = RGBColor(0xB3, 0x26, 0x1E)


def set_size(sh, pt, line=None, italic=None):
    paras = sh.text_frame.paragraphs if line is None else [sh.text_frame.paragraphs[line]]
    for para in paras:
        for r in para.runs:
            r.font.size = Pt(pt)
            if italic is not None:
                r.font.italic = italic


def set_line_geom(sh, x1, y1, x2, y2):
    place(sh, min(x1, x2), min(y1, y2), abs(x2 - x1), abs(y2 - y1))
    xfrm = sh._element.spPr.xfrm
    if x2 < x1:
        xfrm.set("flipH", "1")
    else:
        xfrm.attrib.pop("flipH", None)
    if y2 < y1:
        xfrm.set("flipV", "1")
    else:
        xfrm.attrib.pop("flipV", None)


def append_notes(slide, text):
    tf = slide.notes_slide.notes_text_frame
    tf.text = tf.text + "\n" + text


prs = load()
assert len(prs.slides._sldIdLst) == 28

s_rows_proto = slide_by_title(prs, "July 19, 2024")
proto_row_block = one(s_rows_proto, "Block")
proto_row_caption = one(s_rows_proto, "RowCaption")
proto_row_conn = one(s_rows_proto, "Connector")
s5 = slide_by_title(prs, "Explanation 1: A component-level failure")
proto_block = one(s5, "Block")
proto_card = one(s5, "Card")
proto_conn = one(s5, "Connector")
proto_label = one(s5, "Label")

# =============== P0 — Slide 3: introduce the parts before the failure ===============
s = slide_by_title(prs, "A security update triggers system crashes")
chain_block = one(s, "Block")
chain_cards = sorted(shapes_named(s, "Card"), key=lambda c: c.left)
chain_conns = sorted(shapes_named(s, "Connector"), key=lambda c: c.left)
lab = one(s, "Label")
retext(chain_block, "DEFECTIVE CONTENT")
place(chain_block, 0.90, 4.70, 2.40, 0.70)
retext(chain_cards[0], "sensor crash")
place(chain_cards[0], 3.90, 4.70, 2.20, 0.70)
retext(chain_cards[1], "Windows crash")
place(chain_cards[1], 6.70, 4.70, 2.20, 0.70)
set_line_geom(chain_conns[0], 3.30, 5.05, 3.90, 5.05)
set_line_geom(chain_conns[1], 6.10, 5.05, 6.70, 5.05)
retext(lab, "JULY 19, 2024: ONE RAPID RESPONSE CONTENT UPDATE WAS DEFECTIVE")
place(lab, 0.60, 4.25, 8.80, 0.35)
set_size(lab, 12)
rows = [
    ("CROWDSTRIKE", "a security vendor; organizations install its Falcon product to protect their computers"),
    ("FALCON SENSOR", "software running on each Windows endpoint, with privileged (deep) system access"),
    ("RAPID RESPONSE CONTENT", "detection content the sensor interprets — distributed separately from sensor software"),
]
for (bt, ct), y in zip(rows, [1.20, 2.20, 3.20]):
    b = clone_into(s, proto_row_block)
    retext(b, bt)
    place(b, 0.60, y, 2.80, 0.62)
    set_size(b, 12)
    c = clone_into(s, proto_row_caption)
    retext(c, ct)
    place(c, 3.70, y, 5.70, 0.62)
for y1, y2 in [(1.82, 2.20), (2.82, 3.20)]:
    c = clone_into(s, proto_row_conn)
    set_line_geom(c, 2.00, y1, 2.00, y2)
fix_ids(s)
tf = s.notes_slide.notes_text_frame
tf.text = """TIMING: 3 minutes. ESTABLISH THE SYSTEM BEFORE ANALYZING IT.
ENGINEERING POINT: a student should never have to infer what a component does while being asked to analyze why it failed. So first the parts, in order: CrowdStrike is a security vendor; its Falcon sensor runs on each Windows endpoint with privileged system access (it must see what the operating system sees); the sensor interprets Rapid Response Content — detection content describing what to look for — which is distributed separately from sensor-software releases. THEN the event: on July 19, one content update was defective, the sensor crashed processing it, and because the sensor runs with privileged access, the sensor's crash became the operating system's crash.
DOCUMENTED (CrowdStrike RCA, 6 Aug 2024): the defective artifact was a Rapid Response Content update, not an ordinary Windows OS update; Rapid Response Content could be updated independently of sensor-code releases — a faster release path. HOLD that faster-path observation for Movement 2 (slide 7) unless a student raises it; here it is premature.
ASK (the footer): why can security content crash an operating system?
EXPECTED ANSWERS: 'it's code' — sharpen: it is DATA interpreted by privileged code, and data consumed by privileged code can change its behavior. Misconception: students assume the update was a program; the content/interpreter split is the load-bearing distinction for the whole lecture.
TRANSITION: what exactly did the content and its interpreter disagree about? Numbers next.
SOURCES: CrowdStrike RCA, 6 Aug 2024."""

# =============== P0 — Slide 4: bridge sentence, then the (shifted) diagram ===============
s = slide_by_title(prs, "The content and the interpreter disagreed")
for sh in list(s.shapes):
    if sh.name.split(" ")[0] in ("Block", "Card", "Connector", "Label"):
        delete_shape(sh)
bridge = clone_into(s, proto_label)
retext(bridge,
       "The sensor interprets detection content using predefined templates.",
       "Each template specifies the input fields the interpreter expects.")
place(bridge, 0.60, 0.95, 8.80, 0.62)
set_size(bridge, 13, italic=True)

def blk4(lines, l, t, w, h):
    b = clone_into(s, proto_block); retext(b, *lines); place(b, l, t, w, h); return b

def crd4(lines, l, t, w, h):
    c = clone_into(s, proto_card); retext(c, *lines); place(c, l, t, w, h); return c

def ln4(x1, y1, x2, y2):
    c = clone_into(s, proto_conn); set_line_geom(c, x1, y1, x2, y2); return c

blk4(["TEMPLATE DEFINITION: 21 FIELDS"], 3.00, 1.75, 4.00, 0.55)
crd4(["content expects 21 inputs"], 0.90, 2.65, 3.60, 0.55)
crd4(["integration code supplies 20"], 5.50, 2.65, 3.60, 0.55)
ln4(4.40, 2.30, 2.70, 2.65)
ln4(5.60, 2.30, 7.30, 2.65)
blk4(["INPUT ARRAY [0 … 19]"], 3.40, 3.60, 3.20, 0.55)
ln4(2.70, 3.20, 4.20, 3.60)
ln4(7.30, 3.20, 5.80, 3.60)
crd4(["interpreter attempts access [20]"], 3.20, 4.50, 3.60, 0.55)
ln4(5.00, 4.15, 5.00, 4.50)
oob = blk4(["OUT-OF-BOUNDS READ"], 3.50, 5.35, 3.00, 0.55)
oob.line.color.rgb = RED
oob.line.width = Pt(2.0)
ln4(5.00, 5.05, 5.00, 5.35)
fix_ids(s)
append_notes(s, """EXPOSITION ORDER (author): read the bridge sentence FIRST, verbatim — templates are the interface between detection content and the interpreter; each template specifies the input fields the interpreter expects. Only then walk the diagram; it is now intelligible without inference. Keep the concrete indices [0 … 19] and [20] — they are what makes the failure technically understandable.""")

# =============== P0 — Slide 14: concrete control example before the terminology ===============
s = slide_by_title(prs, "Explanation 3: Safety depends on control")
ex = clone_into(s, proto_label)
retext(ex,
       "A release manager authorizes an update (a control action) based on test results (feedback).",
       "If the tests do not represent actual endpoint behavior, the manager’s picture of the",
       "system (the process model) is wrong — and a well-intended decision is unsafe.")
place(ex, 0.60, 0.95, 8.80, 0.95)
set_size(ex, 13, italic=True)
ctrl = next(b for b in shapes_named(s, "Block") if "CONTROLLER" in b.text_frame.text)
proc = next(b for b in shapes_named(s, "Block") if "CONTROLLED" in b.text_frame.text)
place(ctrl, 3.00, 2.10, 4.00, 1.00)
place(proc, 3.00, 4.20, 4.00, 0.80)
conns = [c for c in shapes_named(s, "Connector") if c.width == 0]
assert len(conns) == 2
for c in conns:
    up = c._element.spPr.xfrm.get("flipV") == "1"
    x = 5.70 if up else 4.30
    if up:
        set_line_geom(c, x, 4.20, x, 3.10)
    else:
        set_line_geom(c, x, 3.10, x, 4.20)
for l in shapes_named(s, "Label"):
    t = l.text_frame.text
    if "control actions" in t:
        place(l, 1.20, 3.35, 2.60, 0.45)
    elif "feedback" in t and "release manager" not in t:
        place(l, 7.10, 3.35, 2.30, 0.45)
fix_ids(s)
append_notes(s, """EXPOSITION ORDER (author): lead with the CONCRETE example at the top of the slide — release manager, test results, wrong picture — and only then introduce the four terms by mapping them onto it: the manager is the CONTROLLER, the authorization is a CONTROL ACTION, the test results are FEEDBACK, the manager's picture of the system is the PROCESS MODEL. The abstraction arrives after the student has an instance of it. Slide 15 then applies the full structure to CrowdStrike — do not merge or reorder the two.""")

# =============== P0 — Slide 26: introduce Cloudflare/WAF, mechanism before measurements ===============
s = slide_by_title(prs, "July 2, 2019: Cloudflare’s WAF outage")
rows = list(zip(shapes_named(s, "Block"), shapes_named(s, "RowCaption")))
assert len(rows) == 4
content = [
    ("CLOUDFLARE", "operates a global edge network; its WAF inspects customers’ web traffic for attacks"),
    ("WAF MANAGED RULES", "attack-detection patterns — often regular expressions — running on every edge server"),
    ("THE DEFECTIVE RULE", ".*(?:.*=.*) — matching work grows explosively with input length (backtracking)"),
    ("27 MINUTES", "CPU exhaustion from 13:42 UTC; rule killed globally 14:07; traffic normal 14:09"),
]
for (b, c), (bt, ct) in zip(rows, content):
    retext(b, bt)
    retext(c, ct)
    set_size(b, 12)
fix_ids(s)
tf = s.notes_slide.notes_text_frame
tf.text = """TIMING: 3 minutes — a transfer exercise, not a second full case study.
EXPOSITION ORDER (author): establish the system first. Cloudflare operates an edge network a large share of web traffic passes through; a Web Application Firewall (WAF) inspects that traffic for attacks; WAF rules describe attack patterns, often as regular expressions. THEN the failure — and explain the MECHANISM before the measurements: some regular expressions require disproportionately more matching work as the input grows, because a backtracking engine re-tries overlapping ways to match. The 23-vs-555 bars at the bottom ILLUSTRATE that mechanism; they do not replace the explanation.
FACTS (Cloudflare post-mortem, 'Details of the Cloudflare outage on July 2, 2019'): a WAF Managed Rule targeting XSS deployed globally at 13:42 UTC — in seconds, with no staged rollout, while software releases went DOG → PIG → Canary; the sub-pattern .*(?:.*=.*) caused catastrophic backtracking (their worked numbers: 'x=x' took 23 steps; 'x=' plus twenty 'x's took 555); PCRE 'has no mechanism to protect against a runaway expression'; WAF identified 14:00, killed globally 14:07, traffic normal 14:09, re-enabled 14:52 — 27 minutes of impact. A CPU protection had been accidentally removed weeks earlier.
DEEP STRUCTURE (lead with it at the footer): in both incidents, content or configuration shipped through a faster, less-gated path than code; and in both, a validator ran and passed while checking the wrong property — arity against the wrong contract at CrowdStrike, runtime instead of CPU at Cloudflare.
CARE: regex matching cost depends on expression, engine, AND input — do not imply all engines behave this way (re2 and Rust's engine carry runtime guarantees). Do not conflate with Cloudflare's November 2025 outage — different mechanism.
ASK: what in this story have you already seen this lecture?
TRANSITION: your turn — diagnose it with the three models."""

# =============== P1 — Slide 8: name the model; show ALIGNMENT ===============
s = slide_by_title(prs, "Explanation 2: Defenses can fail together")
retext(one(s, "Title"), "Explanation 2: The Swiss cheese model")
deflab = next(l for l in shapes_named(s, "Label") if "condition of the defenses" in l.text_frame.text)
retext(deflab,
       "Reason’s claim: every defensive layer can contain weaknesses —",
       "a failure escapes only when the weaknesses ALIGN.")
set_size(deflab, 14)
place(deflab, 0.60, 5.55, 8.80, 0.70)
through = next(c for c in shapes_named(s, "Connector") if Emu(c.width).inches > 5)
through.line.color.rgb = RED
through.line.width = Pt(2.5)
aligned = [h for h in shapes_named(s, "Hole") if abs(Emu(h.top).inches - 3.05) < 0.02]
assert len(aligned) == 4
for h in aligned:
    h.line.color.rgb = RED
    h.line.width = Pt(1.75)
fix_ids(s)
append_notes(s, """NAMING + THE DISTINCTIVE CLAIM (author): say the model's NAME — this is Reason's Swiss cheese model — and state its distinctive claim precisely: every defensive layer can contain weaknesses, and a failure escapes only when weaknesses in successive layers ALIGN. The red path through the red-rimmed holes is the alignment; the other holes are weaknesses that this particular trajectory did not exercise. Without the alignment idea a student leaves with defense-in-depth but not Reason's contribution: defenses fail TOGETHER, as a configuration, not merely one at a time.""")

# =============== P1 — Slide 13: staged reveal in the notes ===============
s = slide_by_title(prs, "What system are we investigating?")
append_notes(s, """STAGING (no animation is authored in this deck — run the reveal verbally): present the boundaries SEQUENTIALLY, never all four at once. Step 1 — cover or ignore the outer boundaries; point at ① content + interpreter and ask 'what caused the crash?'. Step 2 — widen to ② CrowdStrike's engineering system: 'why did the update escape?'. Step 3 — widen to ③ deployment, endpoints, services: 'why did organizations become unavailable?'. Step 4 — widen to ④ the shared-dependency ecosystem: 'why could one failure disrupt many sectors?'. Read each legend line only when its boundary is introduced. The slide is the end-state the staging arrives at.""")

# =============== P1 — Slide 21: supply the missing connection on-slide ===============
s = slide_by_title(prs, "A postmortem is an argument")
lead = clone_into(s, proto_label)
retext(lead, "We can now explain the failure three ways. The next decision: what does the evidence justify CHANGING?")
place(lead, 0.60, 1.00, 8.80, 0.40)
set_size(lead, 13, italic=True)
fix_ids(s)
append_notes(s, """BRIDGE (author): the jump from comparing causal models to reading postmortems needs its connective sentence, now on the slide — once engineers have an explanation, they must determine what the evidence justifies changing. A postmortem is the document where that determination is argued.""")

# =============== P1 — Slide 23: five categories become a sequence of decisions ===============
s = slide_by_title(prs, "Where will the next engineer encounter the lesson?")
blocks = sorted(shapes_named(s, "Block"), key=lambda b: b.top)
assert len(blocks) == 5
for i in range(4):
    y1 = Emu(blocks[i].top).inches + Emu(blocks[i].height).inches
    y2 = Emu(blocks[i + 1].top).inches
    c = clone_into(s, proto_conn)
    set_line_geom(c, 2.40, y1, 2.40, y2)
mapping = clone_into(s, proto_label)
retext(mapping, "obligation → Requirements · representation → Modeling · evidence + evaluation → Validation · enforcement → Alignment")
place(mapping, 0.60, 5.93, 8.80, 0.32)
set_size(mapping, 11)
fix_ids(s)
append_notes(s, """SEQUENCE, NOT TAXONOMY (author): walk the five rows as CONNECTED DECISIONS, each depending on the previous — what obligation was violated? → how should it be REPRESENTED so it can be reasoned about? → what EVIDENCE establishes compliance with it? → how is that evidence EVALUATED? → what ENFORCES that an unacceptable change cannot proceed? The arrows down the left column carry the dependency. Map each to the unit where students met it (the bottom line): obligations in Requirements, representation in Modeling, evidence and evaluation in Validation, enforcement in Alignment. The diagnosis question is where the CHAIN first broke, and the encoding belongs at that link.""")

# =============== P2 — Slide 7: explain wildcard masking ===============
s = slide_by_title(prs, "How did the update reach production?")
lqs = sorted(shapes_named(s, "LevelQuestion"), key=lambda q: q.top)
retext(lqs[1], "passed — every test used a wildcard (‘match anything’) for field 21, so field 21 was never read")
append_notes(s, """WILDCARD, SPELLED OUT (author): a wildcard matching criterion means 'accept any value here'. When the matching criterion for field 21 is a wildcard, the Content Interpreter never needs to READ field 21 to decide the match — so the missing 21st array entry was never touched, and every wildcard test passed honestly. The mismatch was real the whole time; the tests simply never exercised the read that exposes it. That is what makes it a latent condition rather than a flaky test.""")

# =============== P2 — Slide 10: shared timeline for the four interventions ===============
s = slide_by_title(prs, "What should change?")
blocks = sorted(shapes_named(s, "Block"), key=lambda b: b.top)
stage = ["acts at the defect", "acts before release", "acts during propagation", "acts after disruption"]
for b, tag in zip(blocks, stage):
    head = b.text_frame.text.strip().split("\n")[0]
    retext(b, head, tag)
    set_size(b, 9, line=1)
fix_ids(s)
append_notes(s, """TIMELINE (author): give the four a shared timeline — defect → release → propagation → disruption — and read each intervention against the stage it addresses (now tagged on each block): PREVENT acts at the defect, DETECT before release, CONTAIN during propagation, RECOVER after disruption. That turns the taxonomy into an analytical tool: for any proposed intervention, ask which stage it acts on and what has already happened by then.""")

# =============== P2 — Slide 19: authority + evidence, on the slide ===============
s = slide_by_title(prs, "Who can impose a system-wide constraint?")
bot = next(l for l in shapes_named(s, "Label") if "Not every sector" in l.text_frame.text)
retext(bot,
       "Each controller requires both authority to act and evidence about the consequences.",
       "Not every sector has a regulator empowered to mandate these controls.",
       "Dashed control = hypothetical addition to the structure, not a documented arrangement.")
place(bot, 0.60, 5.55, 8.80, 0.85)
set_size(bot, 12)
for r in bot.text_frame.paragraphs[0].runs:
    r.font.bold = True
append_notes(s, """NOT A HIERARCHY (author): the stack can misread as an org chart. The corrective sentence is now on the slide — each controller requires BOTH authority to act AND evidence about the consequences. A regulator with authority but no operational evidence cannot control well; an operator rich in evidence may lack the authority to impose a sector-wide requirement. Control needs the pair, at every level.""")

# =============== P2 — Slide 22: introduce reflective practice first ===============
s = slide_by_title(prs, "Reflection changes the engineer")
intro = clone_into(s, proto_label)
retext(intro,
       "Engineers develop expertise by examining the consequences of their decisions and",
       "incorporating what they learn into future judgment — reflection, not mere accumulation of experience.")
place(intro, 0.60, 0.95, 8.80, 0.62)
set_size(intro, 13, italic=True)
chain = sorted([sh for sh in s.shapes if sh.name.split(" ")[0] in ("Block", "Card")],
               key=lambda x: x.top)
assert len(chain) == 5
tops = [1.62, 2.44, 3.26, 4.08, 4.90]
for sh, t in zip(chain, tops):
    place(sh, 2.85, t, 4.30, 0.60)
conns = sorted([c for c in shapes_named(s, "Connector") if c.width == 0], key=lambda c: c.top)
assert len(conns) == 4
for c, t in zip(conns, [2.22, 3.04, 3.86, 4.68]):
    set_line_geom(c, 5.00, t, 5.00, t + 0.22)
schon = next(l for l in shapes_named(s, "Label") if "Experience alone" in l.text_frame.text)
place(schon, 0.60, 5.60, 8.80, 0.65)
set_size(schon, 13)
fix_ids(s)
append_notes(s, """EXPOSITION ORDER (author): read the intro sentence first, then walk the chain. The distinction to land before the diagram: experience accumulates by itself; expertise does not. Reflection is the deliberate comparison of expectation with outcome — that comparison, not the hours, is what feeds the repertoire.""")

# =============== Slide 20 — two dimensions, not one continuum ===============
s = slide_by_title(prs, "What did each model let us change?")
labels = shapes_named(s, "Label")
by_text = {l.text_frame.text.strip(): l for l in labels}
q_new = ["What mechanism produced the failure?",
         "Why did defenses fail to prevent it?",
         "Which safety constraints were inadequately controlled?"]
i_new = ["repair the defective mechanism",
         "strengthen or diversify defenses",
         "change control actions, feedback, or authority"]
qs = [by_text["What failed?"], by_text["Why did it escape?"], by_text["Why so broad an impact?"]]
ivs = [by_text["repair the defect"], by_text["improve the defenses"],
       by_text["change controls and responsibilities"]]
xs = [0.70, 3.90, 7.10]
for l, x, t in zip(qs, xs, q_new):
    retext(l, t)
    place(l, x, 2.25, 2.20, 0.95)
    set_size(l, 12)
for l, x, t in zip(ivs, xs, i_new):
    retext(l, t)
    place(l, x, 3.35, 2.20, 0.80)
    set_size(l, 12)
delete_shape(one(s, "Connector"))  # the narrow-to-broad axis
for t in ["Narrower", "Broader", "causal boundary"]:
    delete_shape(by_text[t])
big = next(l for l in shapes_named(s, "Label") if "Broader explanations" in l.text_frame.text)
retext(big,
       "Any of the three can be applied at different system boundaries — this lecture chose progressively broader ones.",
       "Broader explanations require more evidence and may require greater authority to act.",
       "Choose the model for its question, and the boundary sufficient for the decision.")
place(big, 0.70, 4.55, 8.60, 1.30)
set_size(big, 14)
fix_ids(s)
tf = s.notes_slide.notes_text_frame
tf.text = """TIMING: 3 minutes. THE SYNTHESIS SLIDE — closes Movement 3. (Backbone slide: models enable different interventions.)
ENGINEERING POINT — TWO DIMENSIONS, NOT ONE: the three models differ in their EXPLANATORY MECHANISM (what kind of account they construct: a causal chain; a configuration of defenses; a control structure enforcing constraints) and, SEPARATELY, in the SYSTEM BOUNDARY they are applied to. This lecture applied them at progressively broader boundaries because the investigation's questions widened — but that was our choice of boundary, not a property of the models.
THE CONFLATION TO PREVENT: STAMP is NOT simply a more comprehensive Swiss cheese. STAMP can be applied to a narrow technical system (a single controller and its process); Swiss cheese can be applied across organizational boundaries (defenses spanning vendor and customer). The model answers its question; the boundary is chosen to fit the decision.
WALK THE TABLE: component-level — what mechanism produced the failure? — repair the defective mechanism. Swiss cheese — why did defenses fail to prevent it? — strengthen or diversify defenses. STAMP — which safety constraints were inadequately controlled? — change control actions, feedback, or authority.
ASK: broader explanations expose more — what do they COST? (More evidence to support the account; often more authority to act on it.)
EXPECTED ANSWERS AND MISCONCEPTION: students reflexively crown the broadest model. Keep the standing caution: broader is not automatically better; choose the boundary sufficient for the decision.
TRANSITION: we have candidate interventions at every scope. Which lessons does the evidence actually justify — and how do we keep them?"""

# =============== Slide 1 notes — engineering, not accident analysis ===============
s1_title = prs.slides[0]
append_notes(s1_title, """FRAME IT AS ENGINEERING, NOT ACCIDENT ANALYSIS (author): this unit is Failure-Aware ENGINEERING. The destination is not that students can explain failures with three frameworks; it is that an organization becomes better at preventing, detecting, containing, and recovering from failure because of what it learned — and that the learning can change exactly two things: the engineer's judgment, through reflection, and the engineering environment, through models, constraints, validation mechanisms, and controls. Say this up front; the three models are instruments toward it, not the point.
BUDGET: the TIMING lines across the deck now sum to 75 minutes, with the slide-6 commitment discussion and the slide-27 transfer exercise protected.""")

# =============== Pacing — rebalance TIMING lines to sum to 75 ===============
TIMING = {1: 1, 2: 2, 3: 3, 4: 4, 5: 2, 6: 4, 7: 3, 8: 2, 9: 3, 10: 2, 11: 2, 12: 1,
          13: 3, 14: 3, 15: 2, 16: 4, 17: 3, 18: 4, 19: 2, 20: 3, 21: 2, 22: 2,
          23: 3, 24: 2, 25: 2, 26: 3, 27: 6, 28: 2}
assert sum(TIMING.values()) == 75
for idx, slide in enumerate(prs.slides, 1):
    tf = slide.notes_slide.notes_text_frame
    new = re.sub(r"TIMING: \d+ minutes", f"TIMING: {TIMING[idx]} minutes", tf.text, count=1)
    assert "TIMING:" in new, idx
    if new != tf.text:
        tf.text = new

save_and_verify(prs, 28)
print("exposition pass applied")
