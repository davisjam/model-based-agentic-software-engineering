"""Phase M45: build Movement 4 (Learning, 5 slides per recommendation 4) and
Movement 5 (Transfer, 3 slides) — 53 -> 61 slides.

Facts from VERIFIED-FACTS-261015.md §A (CrowdStrike) and §B/§C (Cloudflare + shared
deep structure). Correction applied: progressive deployment is evidenced by rollout
limits, detection latency, telemetry, and halt capability — not fault injection.
"""
from __future__ import annotations

import sys

sys.path.insert(0, "book/_design/2-4-redux-build/scripts")
from pptx.util import Pt

from buildlib import (
    clone_into, delete_shape, dup_slide, fix_ids, load, one, place, retext,
    save_and_verify, set_notes, shapes_named, slide_by_title, slide_title,
)


def slide_by_prefix(prs, prefix):
    hits = [s for s in prs.slides if slide_title(s).startswith(prefix)]
    assert len(hits) == 1, f"prefix {prefix!r}: {len(hits)} hits"
    return hits[0]


def set_size(sh, pt, line=None):
    paras = sh.text_frame.paragraphs if line is None else [sh.text_frame.paragraphs[line]]
    for para in paras:
        for r in para.runs:
            r.font.size = Pt(pt)


def set_line_geom(sh, x1, y1, x2, y2):
    place(sh, min(x1, x2), min(y1, y2), abs(x2 - x1), abs(y2 - y1))
    xfrm = sh._element.spPr.xfrm
    if x2 < x1:
        xfrm.set("flipH", "1")
    else:
        xfrm.attrib.pop("flipH", None)
    if y2 < y1:
        xfrm.set("flipV", "1")
    else:
        xfrm.attrib.pop("flipV", None)


prs = load()
assert len(prs.slides._sldIdLst) == 53, "expected 53 slides at phase start"

t_steps = slide_by_prefix(prs, "How do we establish")                  # old 27
t_chain = slide_by_title(prs, "How will the next human engineer or agent know?")  # old 32
t_convert = slide_by_title(prs, "Rule: Convert failure lessons into engineering structure")  # old 33
t_twocol = slide_by_title(prs, "Imagine a race condition. Diagnosis reopens the design space")  # old 29
t_rows = slide_by_title(prs, "July 19, 2024")                          # old 35
t_proto = slide_by_title(prs, "Explanation 1: A component-level failure")  # old 38
proto_conn = one(t_proto, "Connector")
proto_label = one(t_proto, "Label")
proto_footer = one(t_proto, "FooterBand")

# =============== Slide 21 — A postmortem is an argument ===============
s = dup_slide(prs, t_steps)
retext(one(s, "Title"), "A postmortem is an argument")
cards = sorted(shapes_named(s, "Card"), key=lambda c: c.top)
assert len(cards) == 4
content = [
    ("1   OBSERVATION", "Windows systems crashed after the content update"),
    ("2   CAUSAL ACCOUNT", "the interpreter read the 21st entry of a 20-entry input array"),
    ("3   LESSON", "content must be validated against the interpreter’s actual obligations"),
    ("4   INTERVENTION", "an independently evaluated compatibility check before release"),
]
for c, (head, cap) in zip(cards, content):
    retext(c, head, cap)
fb = clone_into(s, proto_footer)
retext(fb, "What evidence supports each step?")
fix_ids(s)
set_notes(s, """TIMING: 3 minutes. OPENS MOVEMENT 4 — what should change?
ENGINEERING POINT: a postmortem is not a timeline; it ADVANCES CLAIMS — about what happened, what caused it, what it teaches, and what should change. Each step down this ladder is a judgment that can be wrong, incomplete, or overgeneralized, and each needs its own evidence.
FACTUAL BACKGROUND: steps 1 and 2 are documented (CrowdStrike RCA, 6 Aug 2024: the out-of-bounds read of a 21st entry in a 20-entry input array). Steps 3 and 4 are the INSTRUCTOR'S proposed generalization and intervention — label them as such in class. Keeping documented findings separate from proposed lessons is the discipline this slide teaches.
ASK: what evidence would support the generalization from ONE malformed update to a broad compatibility invariant? (Evidence that the content/interpreter interface has other mismatches available; evidence that the same release path would ship them.)
EXPECTED ANSWERS AND MISCONCEPTION: students treat the lesson as self-evident once the mechanism is known — 'obviously validate content.' Push: WHICH validation, against WHOSE contract? Slide 9 showed a validator that ran and checked the wrong contract; an unjustified lesson can preserve exactly that error.
TRANSITION: suppose the lesson is sound. Where does it go? Two destinations — the engineer, and the engineering environment. First the engineer.""")

# =============== Slide 22 — Reflection changes the engineer ===============
s = dup_slide(prs, t_chain)
retext(one(s, "Title"), "Reflection changes the engineer")
blocks = {b.text_frame.text.strip(): b for b in shapes_named(s, "Block")}
for name in ["MODEL", "APPLICABILITY", "VALIDATOR", "GATE", "ASK ON SLACK"]:
    delete_shape(blocks[name])
for sh in list(shapes_named(s, "Straight")):
    delete_shape(sh)
chain_blocks = sorted(shapes_named(s, "Block"), key=lambda b: b.top)
chain_cards = sorted(shapes_named(s, "Card"), key=lambda c: c.top)
assert len(chain_blocks) == 2 and len(chain_cards) == 2
retext(chain_blocks[0], "DECISION")
retext(chain_cards[0], "consequence")
retext(chain_blocks[1], "REFLECTION")
retext(chain_cards[1], "repertoire")
extra = clone_into(s, chain_blocks[0])
retext(extra, "FUTURE RECOGNITION")
place(extra, 2.85, 4.80, 4.30, 0.60)
conn = clone_into(s, proto_conn)
set_line_geom(conn, 5.00, 4.50, 5.00, 4.80)
lab = clone_into(s, proto_label)
retext(lab, "Experience alone does not produce expertise — reflection compares what was",
       "expected with what occurred.   (Schön, The Reflective Practitioner, 1983)")
place(lab, 0.60, 5.70, 8.80, 0.80)
fb = clone_into(s, proto_footer)
retext(fb, "What will an experienced engineer notice next time?")
fix_ids(s)
set_notes(s, """TIMING: 3 minutes. PRESERVED FROM THE ORIGINAL LECTURE — Schön's contribution.
ENGINEERING POINT: some lessons change the ENGINEER. Donald Schön's reflective practice: decision → consequence → reflection → repertoire → future recognition. Experience does not automatically produce expertise; the engineer must compare what they expected with what occurred and decide what the difference means.
FACTUAL BACKGROUND / CROWDSTRIKE INSTANCE: an engineer who has worked through this incident may later recognize that DATA-ONLY updates can change program behavior in consequential ways — and ask whether an apparently harmless configuration or content change is actually interpreted by privileged code. That recognition is broader than remembering the particular field-count mismatch; the repertoire stores the pattern, not the instance.
ATTRIBUTION: Schön, The Reflective Practitioner (1983).
CARE: do not present reflection as a substitute for mechanical controls — the next two slides are about exactly the lessons that should NOT depend on anyone's memory.
ASK: what will an experienced engineer notice next time?
EXPECTED ANSWERS: 'check array bounds' (too narrow — that is the instance); 'distrust config changes' (too broad — paralyzing). The calibrated answer: treat content consumed by privileged code as code, with the same skepticism.
TRANSITION: reflection lives in a person. People leave, and agents arrive without the organization's history. Where does the lesson live then?""")

# =============== Slide 23 — Where will the next engineer encounter the lesson? ===============
s = dup_slide(prs, t_convert)
retext(one(s, "Title"), "Where will the next engineer encounter the lesson?")
cards = sorted(shapes_named(s, "Card"), key=lambda c: c.top)
assert len(cards) == 5
for c, txt in zip(cards, [
    "Invariant — content must satisfy the interpreter’s contract",
    "Model — the content–interpreter interface, held in one place",
    "Sensor — compatibility test results",
    "Validator — reject content that violates the contract",
    "Gate — block a release that fails validation",
]):
    retext(c, txt)
chs = sorted(shapes_named(s, "ColumnHeading"), key=lambda c: c.left)
retext(chs[1], "ENCODE IT AS — FOR CROWDSTRIKE")
fix_ids(s)
set_notes(s, """TIMING: 4 minutes. GOVERNANCE CONVERSION — the lesson changes the ENVIRONMENT.
ENGINEERING POINT: organizations cannot rely on every future contributor remembering every incident — and the problem sharpens when work is delegated to agents, which may have substantial implementation capability and no knowledge of the organization's incident history. The five left-hand categories are DIAGNOSTIC QUESTIONS, not five mandatory steps: a failure may expose an obligation nobody articulated, an obligation known but not represented, missing evidence, evidence evaluated against the wrong criterion, or a correct evaluation that nothing enforced.
FACTUAL CARE: do NOT claim CrowdStrike lacked all five. Some controls existed and proved inadequate — the Content Validator RAN (against the wrong contract, RCA finding 4). The right-hand column is a proposed engineering response, not a description of CrowdStrike's actual remediation.
ANIMATION-IN-SPEECH: cover the right column; ask students to propose the artifact for each left-hand diagnosis before revealing.
ASK: for THIS incident, which diagnosis fits best? (Evaluation: a validator existed and evaluated the wrong contract — plus two absent constraints.)
EXPECTED ANSWERS AND MISCONCEPTION: 'add more gates.' The footer is the rule: a failure is not a request for another gate — diagnose what was missing, THEN encode it.
TRANSITION: encoding is not the end of the decision. Different lessons land at different ALTITUDES — and need different authorities.""")

# =============== Slide 24 — Different lessons require different authorities ===============
s = dup_slide(prs, t_twocol)
retext(one(s, "Title"), "Different lessons require different authorities")
blocks = sorted(shapes_named(s, "Block"), key=lambda b: b.top)
cards = sorted(shapes_named(s, "Card"), key=lambda c: c.top)
assert len(blocks) == 4 and len(cards) == 4
rows = [
    ("Content must match its interpreter", "schema + compatibility validator", "VENDOR"),
    ("A defective update must not reach every endpoint", "progressive deployment + halt controls", "VENDOR / ENTERPRISE"),
    ("Critical services must tolerate endpoint failure", "isolation + continuity architecture", "ENTERPRISE"),
    ("Shared dependencies correlate failures", "resilience requirements + oversight", "ENTERPRISE / SECTOR AUTHORITY"),
]
for b, c, (lesson, struct, who) in zip(blocks, cards, rows):
    retext(b, lesson)
    retext(c, struct, who)
    set_size(c, 10, line=1)
chs = sorted(shapes_named(s, "ColumnHeading"), key=lambda c: c.left)
retext(chs[0], "LESSON")
retext(chs[1], "ENGINEERING STRUCTURE — AND WHOSE AUTHORITY")
retext(one(s, "FooterBand"),
       "If a hospital must keep delivering care when its endpoint security fails,",
       "whose requirement is that?")
fix_ids(s)
set_notes(s, """TIMING: 3 minutes. THE LOCAL AND SYSTEMIC ANALYSES MEET.
ENGINEERING POINT: the same incident justifies interventions at several levels, and each lesson's destination follows from the property it preserves and the AUTHORITY required to enforce it. A vendor can improve its contracts, validators, and release controls. An enterprise can invest in isolation, staged deployment where supported, and continuity arrangements. A sector authority may be able to impose resilience obligations across organizations — where its mandate permits (do not imply every sector has one).
ASK (the footer question): if a hospital must continue delivering essential services when its endpoint security software fails — whose requirement is that?
EXPECTED ANSWERS: 'CrowdStrike's.' The recognition to produce: the hospital or health system bears responsibility for its own service continuity, while vendors, platform providers, and potentially regulators control parts of the environment that continuity depends on. Responsibilities must be allocated across the machine AND its environment — this is the Requirements and Specification units, closed into a loop.
MISCONCEPTION: reading the authority column as blame assignment. It is a statement about who CAN act, not who is at fault.
TRANSITION: we now have interventions and owners. One question left, and it is the Validation unit's question: how would we know any of this worked?""")

# =============== Slide 25 — How would we know the system is safer? ===============
s = dup_slide(prs, t_twocol)
retext(one(s, "Title"), "How would we know the system is safer?")
blocks = sorted(shapes_named(s, "Block"), key=lambda b: b.top)
cards = sorted(shapes_named(s, "Card"), key=lambda c: c.top)
assert len(blocks) == 4 and len(cards) == 4
rows = [
    ("Compatibility validator", "negative + boundary-case tests across invalid content classes"),
    ("Progressive deployment", "rollout limits, detection latency, telemetry — and a demonstrated halt"),
    ("Endpoint-failure isolation", "continuity exercised under realistic endpoint failures"),
    ("Independent critical paths", "dependency analysis of the actual alternatives, then a continuity test"),
]
for b, c, (left, right) in zip(blocks, cards, rows):
    retext(b, left)
    retext(c, right)
    set_size(c, 12)
chs = sorted(shapes_named(s, "ColumnHeading"), key=lambda c: c.left)
retext(chs[0], "INTERVENTION")
retext(chs[1], "EVIDENCE THAT IT WORKS")
retext(one(s, "FooterBand"),
       "An untested recovery plan is not evidence that recovery will succeed.",
       "What uncertainty remains?")
fix_ids(s)
set_notes(s, """TIMING: 3 minutes. CLOSES THE CROWDSTRIKE INVESTIGATION.
ENGINEERING POINT: a corrective action is a HYPOTHESIS about how the system will behave differently. Do not equate the existence of a control with evidence of its effectiveness. For each intervention, identify the property it is expected to establish and the evidence that would support it — property, evidence mechanism, judgment: the Validation unit's distinction, applied retrospectively.
ROW-BY-ROW CARE (each evidence row matches ITS intervention): a compatibility validator should reject malformed content violating the interpreter's contract — so the evidence is behavior over invalid and boundary input classes, not a demo on valid content. Progressive deployment should limit exposure before detection and permit a stop — so the evidence is rollout limits, detection latency, telemetry, and a DEMONSTRATED ability to halt distribution (a fault-injection exercise may contribute, but it is not the property being established). Continuity arrangements must be EXERCISED under realistic endpoint failures. Independent critical paths need dependency analysis of the actual alternatives — common-cause analysis reveals dependency structure; it does not by itself demonstrate that services will stay available, hence the continuity test.
ASK: what uncertainty remains after all four?
EXPECTED ANSWERS: 'none.' Correct answer: the next failure class nobody has imagined — which is why containment and recovery stay in the portfolio.
TRANSITION: we built a method on one incident. If the method is real, it should transfer. New incident, five years earlier, different company.""")

# =============== Slide 26 — July 2, 2019: Cloudflare’s WAF outage ===============
s = dup_slide(prs, t_rows)
retext(one(s, "Title"), "July 2, 2019: Cloudflare’s WAF outage")
rows = list(zip(shapes_named(s, "Block"), shapes_named(s, "RowCaption")))
assert len(rows) == 4
content = [
    ("A NEW WAF RULE", "deployed globally at 13:42 UTC — in seconds, with no staged rollout"),
    ("THE REGEX", ".*(?:.*=.*)  —  ‘x=x’ took 23 steps; ‘x=’ plus twenty ‘x’s took 555"),
    ("CPU EXHAUSTION", "PCRE has no mechanism to protect against a runaway expression"),
    ("27 MINUTES", "WAF killed globally at 14:07; traffic normal by 14:09"),
]
for (b, c), (bt, ct) in zip(rows, content):
    retext(b, bt)
    retext(c, ct)
retext(one(s, "FooterBand"),
       "How could one regular expression disrupt a global service?",
       "Configuration shipped through a faster, less-gated path than code — again.")
fix_ids(s)
set_notes(s, """TIMING: 2 minutes — a transfer exercise, not a second full case study.
ENGINEERING POINT: lead with the deep structure, because it is what makes this an APPLICATION of the method rather than a second war story. In both incidents, content or configuration shipped through a faster, less-gated path than code; and in both, a validator ran and passed while checking the wrong property — arity against the wrong contract at CrowdStrike, runtime instead of CPU at Cloudflare.
FACTS (Cloudflare post-mortem, 'Details of the Cloudflare outage on July 2, 2019'): a WAF Managed Rule targeting XSS deployed globally at 13:42 UTC; the sub-pattern .*(?:.*=.*) caused catastrophic backtracking (their worked numbers: 'x=x' 23 steps; 'x=' + twenty x's 555 steps); PCRE 'has no mechanism to protect against a runaway expression'; impact 27 minutes — identified 14:00, killed globally 14:07, traffic normal 14:09, WAF re-enabled 14:52. Software releases went DOG → PIG → Canary; the WAF SOP allowed a non-emergency rule to go global without staged rollout, and a CPU-usage protection had been accidentally removed weeks earlier.
CARE: regex matching cost depends on expression, algorithm, AND input — do not imply all engines behave this way (re2 and Rust's engine carry runtime guarantees). Do not conflate with Cloudflare's November 2025 outage — different mechanism.
ASK: before advancing — what in this story have you already seen this lecture?
TRANSITION: your turn. Diagnose it with the three models.""")

# =============== Slide 27 — Diagnose the Cloudflare outage ===============
s = dup_slide(prs, t_twocol)
retext(one(s, "Title"), "Diagnose the Cloudflare outage")
blocks = sorted(shapes_named(s, "Block"), key=lambda b: b.top)
cards = sorted(shapes_named(s, "Card"), key=lambda c: c.top)
assert len(blocks) == 4 and len(cards) == 4
delete_shape(blocks[3])
delete_shape(cards[3])
rows = [
    ("COMPONENT-LEVEL", "what failed locally?"),
    ("SWISS CHEESE", "why did the defect escape?"),
    ("STAMP", "what controls and feedback permitted the consequences?"),
]
for b, c, (left, right) in zip(sorted(shapes_named(s, "Block"), key=lambda b: b.top),
                               sorted(shapes_named(s, "Card"), key=lambda c: c.top), rows):
    retext(b, left)
    retext(c, right)
chs = sorted(shapes_named(s, "ColumnHeading"), key=lambda c: c.left)
retext(chs[0], "MODEL")
retext(chs[1], "ITS QUESTION")
retext(one(s, "FooterBand"),
       "YOUR TASK — choose an intervention · name the failure class it addresses ·",
       "decide where the lesson should live.",
       "What evidence would establish that your intervention works?")
fix_ids(s)
set_notes(s, """TIMING: 6 minutes — pairs or small groups (5), then debrief (1). THE EXERCISE.
ENGINEERING POINT: students apply the models WITHOUT being led. Each group picks ONE model, proposes an intervention the incident's evidence supports, and answers: what the model preserves, what mechanism it explains, what intervention follows, what evidence would support it, who has authority to implement it.
POSSIBLE ANSWERS (Cloudflare post-mortem, July 2019): COMPONENT-LEVEL — pathological regex under a backtracking engine; reject dangerous complexity, move to an engine with runtime guarantees (Cloudflare's own remediation: re2 or Rust regex), or impose resource bounds. SWISS CHEESE — defenses existed (PR approval, CI/CD, change tickets, SOPs, alerting) and failed in sequence: the test suite measured RUNTIME, not CPU — 'no increase in test suite run time was observed' with the fatal rule; a CPU protection had been accidentally REMOVED weeks earlier; the SOP permitted a non-emergency rule to skip staged rollout. Note the contrast with slide 9: at Cloudflare a defense existed and was removed — a different finding from one never built. STAMP — a constraint like 'a WAF rule update must not consume enough resources to compromise unrelated traffic'; examine rule authors, validation, deployment controllers, monitoring, and what feedback each received.
GRADING STANCE: do not reward choosing the broadest model. Reward an intervention justified by the model and the evidence, addressing a consequential failure class at acceptable cost.
TRANSITION: close the unit — and Act II.""")

# =============== Slide 28 — Failure should change future engineering ===============
s = dup_slide(prs, t_chain)
retext(one(s, "Title"), "Failure should change future engineering")
blocks = {b.text_frame.text.strip(): b for b in shapes_named(s, "Block")}
for name in ["VALIDATOR", "GATE", "ASK ON SLACK"]:
    delete_shape(blocks[name])
for sh in list(shapes_named(s, "Straight")):
    delete_shape(sh)
chain_blocks = sorted(shapes_named(s, "Block"), key=lambda b: b.top)
# after deletions: A NEW ENGINEER (1.2), THE IDEMPOTENCY INVARIANT (3.0), MODEL (5.65), APPLICABILITY (5.9)
retext(chain_blocks[0], "OBSERVE")
chain_cards = sorted(shapes_named(s, "Card"), key=lambda c: c.top)
retext(chain_cards[0], "explain — a causal account the evidence supports")
retext(chain_blocks[1], "LEARN")
delete_shape(chain_cards[1])  # the old "How would they know it exists?" card
# the two branch targets, repurposed from MODEL / APPLICABILITY
branch_l = blocks["MODEL"]
retext(branch_l, "REFLECTION", "changes the engineer")
place(branch_l, 1.30, 4.55, 2.80, 0.85)
set_size(branch_l, 11, line=1)
branch_r = blocks["APPLICABILITY"]
retext(branch_r, "ENGINEERING STRUCTURE", "changes the environment")
place(branch_r, 5.90, 4.55, 2.80, 0.85)
set_size(branch_r, 11, line=1)
out_l = clone_into(s, chain_cards[0])
retext(out_l, "future judgment")
place(out_l, 1.30, 5.80, 2.80, 0.55)
out_r = clone_into(s, chain_cards[0])
retext(out_r, "future constraints")
place(out_r, 5.90, 5.80, 2.80, 0.55)
for pts in [(5.00, 3.60, 2.70, 4.55), (5.00, 3.60, 7.30, 4.55),
            (2.70, 5.40, 2.70, 5.80), (7.30, 5.40, 7.30, 5.80)]:
    c = clone_into(s, proto_conn)
    set_line_geom(c, *pts)
# the old vertical connector below LEARN is replaced by the branch diagonals
from pptx.util import Emu

for c in shapes_named(s, "Connector"):
    if c.width == 0 and 3.55 < Emu(c.top).inches < 4.2 and abs(Emu(c.left).inches - 5.0) < 0.01:
        delete_shape(c)
fb = clone_into(s, proto_footer)
retext(fb, "What will be different when the next engineer or agent acts?")
fix_ids(s)
set_notes(s, """TIMING: 3 minutes. CLOSES THE UNIT AND ACT II. No new framework on this slide.
ENGINEERING POINT: engineering operates under incomplete knowledge; failures expose discrepancies between what engineers believed and what the system did. A causal explanation determines what lessons the evidence supports. Then two distinct outcomes: REFLECTION changes the engineer — recognition and judgment applied to future situations; ENGINEERING STRUCTURE changes the environment — an invariant, architectural constraint, validator, or release gate preserves the lesson for future work, including work done by agents that never saw the incident. The engineer decides which lessons deserve durable representation and which interventions the evidence justifies.
ACT II CLOSE: Delegation determined what work and authority to grant. Modeling made consequential knowledge explicit. Alignment connected obligations to mechanisms that govern work. Failure-aware engineering uses observed consequences to TEST and REVISE those decisions — the loop closes.
ASK (the closing question, and leave it standing): what will be different when the next engineer or agent acts?
EXPECTED ANSWERS: let several stand without adjudicating — the point is that 'nothing would be different' is the only wrong answer.
CARE: do not add a second aphorism beneath the question. End here.
TRANSITION: none — this is the end of the unit.""")

save_and_verify(prs, 61)
print("Movements 4+5 built: slides 54-61 appended")
