"""Shared helpers for the 2-4 redux deck build.

Discipline (learned the hard way, see VERIFIED-FACTS §"Package-integrity rule"):
- locate-and-assert-all FIRST, then mutate, then save, then verify;
- slide deletion drops the sldId AND the relationship;
- every phase ends with: zero duplicate zip entries, clean zip CRC, expected slide count.
"""
from __future__ import annotations

import copy
import zipfile

from pptx import Presentation
from pptx.util import Inches, Pt

WIP = "book/_design/2-4-redux-build/2-4-redux-WIP.pptx"


def load() -> Presentation:
    return Presentation(WIP)


def slide_title(slide) -> str:
    for sh in slide.shapes:
        if sh.name.startswith("Title") and sh.has_text_frame:
            return sh.text_frame.text.strip()
    return ""


def slide_by_title(prs, title: str):
    hits = [s for s in prs.slides if slide_title(s) == title]
    assert len(hits) == 1, f"title {title!r}: {len(hits)} hits"
    return hits[0]


def shapes_named(slide, name: str) -> list:
    return [sh for sh in slide.shapes if sh.name == name or sh.name.split(" ")[0] == name]


def one(slide, name: str, idx: int = 0):
    hits = shapes_named(slide, name)
    assert len(hits) > idx, f"{name}[{idx}]: only {len(hits)} on slide {slide_title(slide)!r}"
    return hits[idx]


def dup_slide(prs, src):
    """Append a structural copy of src (shapes only; fresh notes). Returns the new slide."""
    dst = prs.slides.add_slide(src.slide_layout)
    for sh in list(dst.shapes):
        dst.shapes._spTree.remove(sh._element)
    for sh in src.shapes:
        dst.shapes._spTree.append(copy.deepcopy(sh._element))
    return dst


def clone_into(slide, shape):
    """Deep-copy a shape element (possibly from another slide) into slide; returns new shape."""
    el = copy.deepcopy(shape._element)
    slide.shapes._spTree.append(el)
    for sh in slide.shapes:
        if sh._element is el:
            return sh
    raise AssertionError("clone not found after append")


def delete_shape(shape) -> None:
    shape._element.getparent().remove(shape._element)


def retext(shape, *lines: str) -> None:
    """Write lines into the shape, one per paragraph, inheriting run-0 formatting per para.

    Extra paragraphs are deleted; missing ones are cloned from the last existing paragraph.
    """
    tf = shape.text_frame
    while len(tf.paragraphs) > len(lines):
        p = tf.paragraphs[-1]._p
        p.getparent().remove(p)
    while len(tf.paragraphs) < len(lines):
        last = tf.paragraphs[-1]._p
        last.addnext(copy.deepcopy(last))
    for para, text in zip(tf.paragraphs, lines):
        runs = para.runs
        if not runs:
            run = para.add_run()
            run.text = text
            continue
        runs[0].text = text
        for r in runs[1:]:
            r._r.getparent().remove(r._r)


def place(shape, left=None, top=None, width=None, height=None) -> None:
    if left is not None:
        shape.left = Inches(left)
    if top is not None:
        shape.top = Inches(top)
    if width is not None:
        shape.width = Inches(width)
    if height is not None:
        shape.height = Inches(height)


def _set_id(sh, new_id: int) -> None:
    el = sh._element
    for child in el.iter():
        if child.tag.endswith("}cNvPr"):
            child.set("id", str(new_id))
            return
    raise AssertionError("no cNvPr")


def fix_ids(slide) -> None:
    seen: set[int] = set()
    nxt = 1000
    for sh in slide.shapes:
        sid = sh.shape_id
        if sid in seen:
            while nxt in seen:
                nxt += 1
            _set_id(sh, nxt)
            seen.add(nxt)
            nxt += 1
        else:
            seen.add(sid)


def set_notes(slide, text: str) -> None:
    slide.notes_slide.notes_text_frame.text = text


def delete_slide(prs, slide) -> None:
    """Remove sldId AND drop the relationship (duplicate-zip-entry defense)."""
    sldIdLst = prs.slides._sldIdLst
    rid_attr = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id"
    for e in list(sldIdLst):
        rid = e.get(rid_attr)
        if prs.part.rels[rid].target_part is slide.part:
            sldIdLst.remove(e)
            prs.part.drop_rel(rid)
            return
    raise AssertionError("slide not found in sldIdLst")


def save_and_verify(prs, expected_count: int) -> None:
    prs.save(WIP)
    with zipfile.ZipFile(WIP) as zf:
        names = zf.namelist()
        assert len(names) == len(set(names)), "DUPLICATE ZIP ENTRIES"
        assert zf.testzip() is None, "zip CRC failure"
    p2 = Presentation(WIP)
    got = len(p2.slides._sldIdLst)
    assert got == expected_count, f"slide count {got} != {expected_count}"
    print(f"VERIFIED: {expected_count} slides, zip clean, no duplicate entries")
