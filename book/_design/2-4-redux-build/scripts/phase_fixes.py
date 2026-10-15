"""Phase FIXES: visual-QA corrections from the rendered-slide inspection pass.

Fixes: wrapped DEPLOYMENT label (s8); unmotivated red accent on row 4 (s9, s10);
boundary-tag pileup + struck-through label (s13); clipped bottom labels
(s14, s17, s19, s22); label-over-block (s15); footer position + constraint size (s16);
intervention-row/axis collision (s20).
"""
from __future__ import annotations

import copy
import sys

sys.path.insert(0, "book/_design/2-4-redux-build/scripts")
from pptx.oxml.ns import qn
from pptx.util import Pt

from buildlib import (
    load, one, place, retext, save_and_verify, shapes_named, slide_by_title,
)


def set_size(sh, pt):
    for para in sh.text_frame.paragraphs:
        for r in para.runs:
            r.font.size = Pt(pt)


def match_fill(src, dst):
    """Make dst's explicit fill identical to src's (both have spPr overrides)."""
    s_sp, d_sp = src._element.spPr, dst._element.spPr
    fills = ["a:noFill", "a:solidFill", "a:gradFill", "a:blipFill", "a:pattFill"]
    src_fill = next((f for t in fills for f in s_sp.findall(qn(t))), None)
    for t in fills:
        for el in d_sp.findall(qn(t)):
            d_sp.remove(el)
    if src_fill is not None:
        # insert after ln would be wrong order; fill precedes a:ln in CT_ShapeProperties
        ln = d_sp.find(qn("a:ln"))
        el = copy.deepcopy(src_fill)
        if ln is not None:
            ln.addprevious(el)
        else:
            d_sp.append(el)


prs = load()
assert len(prs.slides._sldIdLst) == 28

# ---- s8: DEPLOYMENT label wraps at 1.6" width ----
s = slide_by_title(prs, "Explanation 2: Defenses can fail together")
dep = next(l for l in shapes_named(s, "Label") if l.text_frame.text.strip() == "DEPLOYMENT")
place(dep, 6.60, width=1.95)

# ---- s9 + s10: row-4 block inherited the template's red accent ----
for title in ["The investigation identifies weaknesses", "What should change?"]:
    s = slide_by_title(prs, title)
    blocks = sorted(shapes_named(s, "Block"), key=lambda b: b.top)
    assert len(blocks) == 4
    match_fill(blocks[0], blocks[3])

# ---- s13: spread boundary tags; unstrike the "thousands" label ----
s = slide_by_title(prs, "What system are we investigating?")
tags = {l.text_frame.text.strip(): l for l in shapes_named(s, "Label")
        if l.text_frame.text.strip() in "①②③④"}
assert len(tags) == 4
place(tags["③"], 8.90, 1.16)   # top-right of B3
place(tags["④"], 0.58, 1.03)
place(tags["②"], 0.93, 1.32)
thous = next(l for l in shapes_named(s, "Label") if "thousands" in l.text_frame.text)
place(thous, 5.45, 4.84, 4.00, 0.28)

# ---- s14: bottom label clipped under the footer ----
s = slide_by_title(prs, "Explanation 3: Safety depends on control")
lab = next(l for l in shapes_named(s, "Label") if "STAMP" in l.text_frame.text)
set_size(lab, 14)
place(lab, 0.60, 5.45, 8.80, 0.80)

# ---- s15: left label overlaps the UPDATE DISTRIBUTION block ----
s = slide_by_title(prs, "Who controlled the update?")
lab = next(l for l in shapes_named(s, "Label") if "release" in l.text_frame.text)
place(lab, 0.38, 3.05, 2.35, 0.80)
set_size(lab, 14)

# ---- s16: constraint smaller; footer back to the house position ----
s = slide_by_title(prs, "The vendor is not the whole system")
lab = next(l for l in shapes_named(s, "Label") if "SAFETY CONSTRAINT" in l.text_frame.text)
set_size(lab, 14)
place(lab, 0.60, 5.82, 8.80, 0.56)
fb = one(s, "FooterBand")
place(fb, 0.30, 6.40, 9.40, 0.62)

# ---- s17: mid label wrapped into the footer ----
s = slide_by_title(prs, "Many computers do not imply independent failure")
lab = next(l for l in shapes_named(s, "Label") if "Redundancy" in l.text_frame.text)
retext(lab, "Redundancy did not help — every replica shared the same dependency.")
set_size(lab, 14)
place(lab, 0.60, 5.98, 8.80, 0.35)

# ---- s19: bottom label wrapped into the footer ----
s = slide_by_title(prs, "Who can impose a system-wide constraint?")
lab = next(l for l in shapes_named(s, "Label") if "Not every sector" in l.text_frame.text)
set_size(lab, 14)
place(lab, 0.60, 5.95, 8.80, 0.38)

# ---- s20: intervention row collides with the axis ----
s = slide_by_title(prs, "What did each model let us change?")
labels = shapes_named(s, "Label")
by_text = {l.text_frame.text.strip(): l for l in labels}
stamp_iv = by_text["change controls, dependencies, responsibilities"]
retext(stamp_iv, "change controls and responsibilities")
for txt in ["repair the defect", "improve the defenses", "change controls and responsibilities"]:
    l = next(x for x in shapes_named(s, "Label") if x.text_frame.text.strip() == txt)
    set_size(l, 13)
    place(l, top=2.95, height=0.60)
for txt, x in [("Narrower", 0.49), ("Broader", 7.49), ("causal boundary", 3.83)]:
    place(by_text[txt], x, 4.00)
place(one(s, "Connector"), 0.49, 4.45)
big = next(l for l in shapes_named(s, "Label") if "Broader explanations" in l.text_frame.text)
place(big, 0.93, 4.95, 8.20, 0.90)
set_size(big, 15)

# ---- s22: Schön label wrapped into the footer ----
s = slide_by_title(prs, "Reflection changes the engineer")
lab = next(l for l in shapes_named(s, "Label") if "Experience alone" in l.text_frame.text)
set_size(lab, 14)
place(lab, 0.60, 5.68, 8.80, 0.65)

save_and_verify(prs, 28)
print("fixes applied")
